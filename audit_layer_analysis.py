"""Audit native formula/aggregate controls and actual SVG/XLSX files. Flow: README.md."""
from copy import deepcopy
from pathlib import Path
import json
import sys
import time
from unittest.mock import patch
from xml.etree import ElementTree as ET

from PySide6 import QtCore, QtGui, QtWidgets, QtSvg
from openpyxl import load_workbook
from attack_catalog import external_id
from attack_matrix import MATRIX_SLUG, MATRIX_STYLE
from layer_analysis import calculate_layer, ScoreExpression
from layer_analysis_dialogs import FormulaDialog, MatrixExportDialog
from main import MainWindow, configure_application
from matrix_exports import export_matrix, matrix_scene, SVG_NS
from navigator_dialogs import LayerSettingsDialog
from navigator_integration import document_layer, validate_layer, update_annotation
from navigator_visuals import LayerAppearance, DEFAULT_GRADIENT, gradient_color
from ttp_documents import TTPDocument, document_payload

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'data' / 'attack'
OUT = ROOT / 'artifacts'
REPORT = OUT / 'layer-analysis-audit.json'
SVG = OUT / 'layer-analysis-example.svg'
EXCEL = OUT / 'layer-analysis-example.xlsx'
PREVIEW = OUT / 'layer-analysis-svg-preview.png'
UI_SCREEN = OUT / 'layer-analysis-preview.png'
FORMULA_SCREEN = OUT / 'layer-formula-preview.png'
FULL_SVG = OUT / 'layer-analysis-full.svg'
FULL_EXCEL = OUT / 'layer-analysis-full.xlsx'
checks = []


def log(message, current=0, total=0):
    print(message + (f' ({current}/{total})' if total else ''), flush=True)


def check(condition, label):
    if not condition:
        raise AssertionError(label)
    checks.append(label)


def wait(condition, label):
    end = time.monotonic() + 90
    while not condition():
        if time.monotonic() > end:
            raise TimeoutError(label)
        QtWidgets.QApplication.processEvents()
        time.sleep(0.01)


def expect_error(task, label):
    try:
        task()
    except (ValueError, PermissionError, InterruptedError):
        check(True, label)
    else:
        raise AssertionError(label)


def cell(matrix, key, slug='execution'):
    return next(item for item in matrix.items_by_id[key] if item.data(0, MATRIX_SLUG) == slug)


def main():
    log(f'Input file: {DATA / "latest.json"}')
    for path in (REPORT, SVG, EXCEL, PREVIEW, UI_SCREEN, FORMULA_SCREEN, FULL_SVG, FULL_EXCEL):
        log(f'Output file: {path}')
    pointers = {path: path.read_bytes() for path in (DATA / 'active.json', DATA / 'latest.json') if path.exists()}
    log('[Analysis 1/7] Native workspace and safe arithmetic', 0, 7)
    app = QtWidgets.QApplication(sys.argv[:1])
    configure_application(app)
    window = MainWindow(DATA, DATA / 'latest.json')
    window.show()
    wait(lambda: window.loaded_ok and not window.loader.isRunning(), 'startup')
    catalog = window.catalog
    lookup = {external_id(obj): key for key, obj in catalog.records.items()}
    parent, task, at = lookup['T1053'], lookup['T1053.005'], lookup['T1053.002']
    for expression, expected in (('a+b', 100), ('(a+b)/2', 50), ('a-b', -20), ('-a+2*b', 80), ('0.25*a', 10)):
        check(ScoreExpression(expression, {'a', 'b'}).evaluate({'a': 40, 'b': 60}) == expected, expression + ': correct arithmetic')
    for expression in ('__import__("os")', 'a.__class__', 'a[0]', 'a**2', 'a//2', 'True', 'unknown+1', '1e309',
                       '9' * 320, '+'.join(['a'] * 100)):
        expect_error(lambda expr=expression: ScoreExpression(expr, {'a', 'b'}), 'reject unsupported/unsafe expression ' + expression[:40])
    expect_error(lambda: ScoreExpression('a/(b-b)', {'a', 'b'}).evaluate({'a': 1, 'b': 1}), 'division by zero rejected')
    first = window.current_document
    first.name = 'Assessment A'
    first.selected = {parent, task, at}
    first.navigator_layer = {'name': first.name, 'domain': catalog.domain, 'versions': {'attack': catalog.version, 'layer': '4.5', 'navigator': '5.2.0'},
        'gradient': deepcopy(DEFAULT_GRADIENT), 'techniques': [
            {'techniqueID': 'T1053', 'score': 20}, {'techniqueID': 'T1053.005', 'score': 40, 'color': '#ff0000', 'comment': 'Source A'},
            {'techniqueID': 'T1053.002', 'score': 0}]}
    second = TTPDocument('Assessment B', catalog, selected={parent, task, at}, navigator_layer={
        'name': 'Assessment B', 'domain': catalog.domain, 'versions': {'attack': catalog.version, 'layer': '4.5', 'navigator': '5.2.0'},
        'techniques': [{'techniqueID': 'T1053', 'score': 40}, {'techniqueID': 'T1053.005', 'score': 60},
                       {'techniqueID': 'T1053.005', 'tactic': 'execution', 'score': 80}]})
    window.selected = set(first.selected)
    window.refresh_selection()
    source_state = first.fingerprint(), second.fingerprint()
    log('[Analysis 2/7] Per-tactic formula layers, missing scores and source isolation', 1, 7)
    result = calculate_layer([first, second], '(a+b)/2', 'Average result', notify=log)
    style = LayerAppearance(document_layer(result))
    check(style.cell('T1053.005', 'execution')['score'] == 60 and style.cell('T1053.005', 'persistence')['score'] == 50,
          'formula preserves execution-specific source scores without bleeding across tactics')
    check(at not in result.selected and all(row['enabled'] is False for row in result.navigator_layer['techniques'] if row['techniqueID'] == 'T1053.002'),
          'missing-score default remains explicitly unscored and unselected')
    zero = calculate_layer([first, second], 'a+b', 'Zero policy', 'zero', log)
    check(at in zero.selected and LayerAppearance(document_layer(zero)).cell('T1053.002', 'execution')['score'] == 0,
          'explicit zero policy computes missing scores while preserving genuine score zero')
    single = calculate_layer([first, second], 'a', 'Single source', notify=log)
    check(at in single.selected, 'unused variable does not introduce missing scores')
    constant = calculate_layer([first, second], '50', 'Constant', notify=log)
    check(all(row['score'] == 50 for row in constant.navigator_layer['techniques']), 'constant formula applies to selected-source union')
    check(not any(row.get('color') for row in result.navigator_layer['techniques']), 'result gradient is not overridden by inherited source colors')
    check(any(row['name'] == 'Score expression' for row in result.navigator_layer['metadata']), 'formula and source identities recorded as provenance')
    expect_error(lambda: calculate_layer([first, second], 'a/(b-b)', 'Invalid', notify=log), 'bad cell arithmetic rejects complete result atomically')
    foreign = TTPDocument('Foreign', catalog.for_domain('ics-attack'))
    expect_error(lambda: calculate_layer([first, foreign], 'a+b', 'Invalid', notify=log), 'cross-domain calculation rejected')
    check((first.fingerprint(), second.fingerprint()) == source_state, 'all formula calculations preserve source documents')

    log('[Analysis 3/7] Average/max/min/sum, zero counting and native settings', 2, 7)
    for operation, expected in (('average', 20), ('max', 40), ('min', 0), ('sum', 60)):
        probe = document_layer(first)
        probe['layout'] = {'showAggregateScores': True, 'aggregateFunction': operation, 'countUnscored': False}
        appearance = LayerAppearance(probe, catalog)
        check(appearance.cell('T1053', 'execution')['aggregateScore'] == expected, operation + ': includes parent and scored children including score zero')
        check(appearance.cell('T1053', 'execution')['score'] == 20, operation + ': original parent score preserved')
        check(appearance.cell('T1053', 'execution')['background'] == gradient_color(expected, DEFAULT_GRADIENT), operation + ': aggregate controls parent gradient')
    probe['layout'].update(aggregateFunction='average', countUnscored=True)
    appearance = LayerAppearance(probe, catalog)
    check(abs(appearance.cell('T1053', 'execution')['aggregateScore'] - 60 / len(appearance.families[('T1053', 'execution')])) < 1e-12,
          'count unscored as zero includes mapped enabled family members in denominator')
    probe['layout']['aggregateFunction'] = 'min'
    probe['techniques'] = [{'techniqueID': 'T1053', 'score': 20}, {'techniqueID': 'T1053.005', 'score': 40}]
    check(LayerAppearance(probe, catalog).cell('T1053', 'execution')['aggregateScore'] == 0, 'unscored-zero policy participates in minimum function')
    probe['techniques'].append({'techniqueID': 'T1053.005', 'tactic': 'persistence', 'score': 99, 'enabled': False})
    probe['layout'].update(countUnscored=False, aggregateFunction='sum')
    check(LayerAppearance(probe, catalog).cell('T1053', 'persistence')['aggregateScore'] == 20, 'disabled child scope excluded from aggregate')
    invalid = deepcopy(probe)
    invalid['layout']['aggregateFunction'] = 'unknown'
    expect_error(lambda: validate_layer(invalid), 'unknown aggregation rejected before UI commit')
    probe['techniques'] = [{'techniqueID': 'T1053.005', 'color': '#112233'}]
    check(LayerAppearance(probe, catalog).cell('T1053', 'execution')['aggregateScore'] is None, 'fully unscored family stays unscored')

    window.install_imported_document(second)
    wait(lambda: window.loaded_ok, 'second source')
    source_state = first.fingerprint(), second.fingerprint()
    window.formula_action.trigger()
    dialog = window.integration_dialog
    check(isinstance(dialog, FormulaDialog) and dialog.sources_table.rowCount() == 2, 'Navigator menu opens native variable/source mapping')
    dialog.expression.setText('a+b')
    dialog.name.setText('Computed & research')
    dialog.missing.setCurrentIndex(dialog.missing.findData('zero'))
    QtWidgets.QApplication.processEvents()
    check(dialog.grab().save(str(FORMULA_SCREEN)), 'formula dialog preview saved')
    dialog.create_button.click()
    wait(lambda: dialog.result_document is not None and not dialog.worker.isRunning() and window.loaded_ok, 'native formula')
    calculated = window.current_document
    check(calculated is dialog.result_document and calculated.selected == {parent, task, at}, 'background formula installs independent result tab')
    check((first.fingerprint(), second.fingerprint()) == source_state, 'native formula leaves all source state unchanged')
    window.open_matrix()
    matrix = window.matrix_view
    wait(lambda: not matrix.building, 'calculated matrix')
    errors = []
    def save_settings():
        settings = app.activeModalWidget()
        try:
            check(isinstance(settings, LayerSettingsDialog), 'native aggregate controls in layer settings')
            settings.show_aggregate.setChecked(True)
            settings.aggregate_function.setCurrentIndex(settings.aggregate_function.findData('sum'))
            settings.maximum.setText('200')
            settings.save()
            check(settings.result() == QtWidgets.QDialog.DialogCode.Accepted, 'aggregate settings accepted atomically')
        except Exception as error:
            errors.append(error)
            settings.reject()
    QtCore.QTimer.singleShot(60, save_settings)
    window.layer_settings_action.trigger()
    if errors:
        raise errors[0]
    check(cell(matrix, parent).data(0, MATRIX_STYLE)['aggregateScore'] == 180, 'native parent displays computed family sum')
    check('Aggregate score: 180' in cell(matrix, parent).toolTip(0), 'native aggregate tooltip identifies derived score')
    check(document_payload(calculated)['navigator_layer']['layout']['aggregateFunction'] == 'sum', 'aggregate settings persist in Studio JSON')
    note = '=SUM(1,2)\n中文 <tag> & details'
    update_annotation(calculated, task, 'execution', note, 120, '#abcdef')
    window.refresh_selection()
    # This export fixture intentionally checks labels containing both identities and names.
    window.change_matrix_display({'showID':True,'showName':True})
    matrix.platform.setCurrentText('Windows')
    wait(lambda: not matrix.building, 'Windows export scope')
    matrix.search.setText('Scheduled Task/Job')
    matrix.expand_button.click()
    state = calculated.fingerprint(), calculated.name, calculated.export_path, calculated.exported_fingerprint

    log('[Analysis 4/7] Native background SVG export and actual vector pixels', 3, 7)
    export_dialog = window.export_matrix_view('svg', SVG)
    check(isinstance(export_dialog, MatrixExportDialog), 'native SVG export dialog')
    export_dialog.export_button.click()
    wait(lambda: export_dialog.output_path is not None and not export_dialog.worker.isRunning(), 'SVG export')
    export_dialog.reject()
    log(f'Input file: {SVG}')
    tree = ET.parse(SVG)
    namespace = {'s': SVG_NS}
    check(tree.getroot().tag == '{' + SVG_NS + '}svg', 'output parses as SVG vector document')
    texts = [node.text or '' for node in tree.findall('.//s:text', namespace)]
    check('Computed & research' in texts, 'XML-escaped layer name survives readback')
    check(any('Aggregate: 180' in text for text in texts), 'SVG includes aggregate score text')
    titles = tree.findall('.//s:rect/s:title', namespace)
    check(any(note in (title.text or '') for title in titles), 'SVG safely embeds literal multiline annotation context')
    colored = next(rect for rect in tree.findall('.//s:rect', namespace) if rect.get('fill') == '#abcdef')
    renderer = QtSvg.QSvgRenderer(str(SVG))
    check(renderer.isValid(), 'SVG accepted by actual Qt vector renderer')
    size = renderer.defaultSize()
    image = QtGui.QImage(size, QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtCore.Qt.GlobalColor.transparent)
    painter = QtGui.QPainter(image)
    renderer.render(painter)
    painter.end()
    check(image.pixelColor(int(colored.get('x')) + int(colored.get('width')) - 5, int(colored.get('y')) + 5).name() == '#abcdef',
          'actual exported SVG paints manual cell color correctly')
    check(image.save(str(PREVIEW)), 'SVG raster preview saved for visual inspection')

    log('[Analysis 5/7] Native colored Excel workbook and note readback', 4, 7)
    export_dialog = window.export_matrix_view('xlsx', EXCEL)
    export_dialog.export_button.click()
    wait(lambda: export_dialog.output_path is not None and not export_dialog.worker.isRunning(), 'Excel export')
    export_dialog.reject()
    log(f'Input file: {EXCEL}')
    book = load_workbook(EXCEL)
    check(book.sheetnames == ['Matrix', 'Techniques', 'Layer Info'], 'Excel includes colored matrix, detailed values and provenance sheets')
    cells = [item for row in book['Matrix'] for item in row if isinstance(item.value, str) and item.value.startswith('> T1053.005')]
    execution_cell = next(item for item in cells if item.fill.fgColor.rgb == '00ABCDEF')
    check(note in execution_cell.comment.text, 'Excel matrix retains complete Unicode/newline note as a comment')
    check(execution_cell.font.color.rgb in ('00000000', '00FFFFFF'), 'Excel colored cells use readable contrast')
    details = list(book['Techniques'].values)
    entry = next(row for row in details[1:] if row[0] == 'Execution' and row[1] == 'T1053.005')
    check(entry[6] == 120 and entry[9] == note, 'Excel detail sheet retains numeric score and literal note text')
    aggregate_row = next(row for row in details[1:] if row[0] == 'Execution' and row[1] == 'T1053')
    check(aggregate_row[6] == 60 and aggregate_row[7] == 180, 'Excel distinguishes original parent score from aggregate')
    check(all(item.data_type != 'f' for sheet in book for row in sheet for item in row), 'user-provided names/comments are literal Excel text')
    check(book['Matrix'].freeze_panes == 'A4' and book['Techniques'].auto_filter.ref, 'Excel matrix/header navigation and detail filtering configured')
    book.close()
    check((calculated.fingerprint(), calculated.name, calculated.export_path, calculated.exported_fingerprint) == state,
          'SVG/Excel artifacts preserve document name, annotation data and save state')
    export_matrix(calculated, FULL_SVG, 'svg', log)
    export_matrix(calculated, FULL_EXCEL, 'xlsx', log)
    log(f'Input file: {FULL_SVG}')
    full_svg = ET.parse(FULL_SVG).getroot()
    check(len(full_svg.findall('.//s:rect/s:title', namespace)) > 700, 'full SVG includes techniques outside filtered/visible viewport')
    check(all(float(rect.get('y')) + float(rect.get('height')) <= float(full_svg.get('height'))
              for rect in full_svg.findall('.//s:rect', namespace)), 'full SVG canvas includes every rendered cell without vertical cropping')
    log(f'Input file: {FULL_EXCEL}')
    full_book = load_workbook(FULL_EXCEL)
    check(full_book['Matrix'].max_column == len(catalog.tactic_order), 'full Excel includes all tactic columns outside horizontal viewport')
    check(full_book['Techniques'].max_row > 700, 'full Excel includes complete matrix technique/tactic rows')
    full_book.close()

    log('[Analysis 6/7] Overlay/filter projection, failure atomicity and domain support', 5, 7)
    scene = matrix_scene(calculated, 'Windows', 'Scheduled Task/Job')
    check(all(cell['id'].startswith('T1053') for col in scene['columns'] for cell in col['cells']), 'matrix exports honor literal name/ID filter')
    expected = next(cell for col in scene['columns'] if col['slug'] == 'execution' for cell in col['cells'] if cell['id'] == 'T1053.005')
    check(expected['background'] == cell(matrix, task).background(0).color().name(), 'export colors match current native matrix')
    overlay = matrix_scene(calculated, 'Windows', 'Scheduled Task/Job', sources=[first, second, calculated])
    check(overlay['overlay'] and any(cell['style'].get('sources') for col in overlay['columns'] for cell in col['cells']), 'export projection includes compatible-layer blending and provenance')
    neutral = matrix_scene(calculated, 'Windows', 'Scheduled Task/Job', colors=False)
    check(all(cell['background'] == '#222526' for col in neutral['columns'] for cell in col['cells']), 'hide-colors view honored by export')
    protected = SVG.read_bytes()
    with patch.object(Path, 'replace', side_effect=PermissionError('audit publication blocked')):
        expect_error(lambda: export_matrix(calculated, SVG, 'svg', log, 'Windows', 'Scheduled Task/Job'), 'failed SVG publish reported')
    check(SVG.read_bytes() == protected and not list(OUT.glob(SVG.name + '.*.tmp')), 'failed SVG publish preserves prior file and cleans temporary output')
    def cancel(message, current=0, total=0):
        if '[Matrix Export 3/4]' in message:
            raise InterruptedError('audit cancelled before write')
    expect_error(lambda: export_matrix(calculated, SVG, 'svg', cancel, 'Windows', 'Scheduled Task/Job'), 'export cancellation stops before publication')
    check(SVG.read_bytes() == protected, 'cancelled export preserves prior file')
    expect_error(lambda: export_matrix(calculated, SVG, 'svg', log, query='no such technique'), 'empty filtered export rejected')
    for domain in ('mobile-attack', 'ics-attack'):
        domain_catalog = catalog.for_domain(domain)
        key = next(iter(domain_catalog.records))
        domain_doc = TTPDocument('Domain test', domain_catalog, selected={key})
        update_annotation(domain_doc, key, '', 'Domain note', 0.25, '#12ab34')
        result = calculate_layer([domain_doc], 'a*2', 'Domain result', notify=log)
        check(all(row['score'] == 0.5 for row in result.navigator_layer['techniques']), domain + ': decimal formula result')
        scene = matrix_scene(domain_doc)
        check(any(cell['background'] == '#12ab34' for col in scene['columns'] for cell in col['cells']), domain + ': domain-specific matrix rendering')
    check(all(path.read_bytes() == data for path, data in pointers.items()), 'dataset pointers unchanged')
    window.sidebar.hide()
    window.output_panel.hide()
    QtWidgets.QApplication.processEvents()
    check(window.grab().save(str(UI_SCREEN)), 'native formula/aggregate matrix screenshot saved')
    window.confirm_studio_changes = lambda *_: "discard"
    window.close()
    log('[Analysis 7/7] Publishing report', 6, 7)
    REPORT.write_text(json.dumps({'status': 'PASS', 'checks': checks, 'svg': str(SVG), 'excel': str(EXCEL),
                                 'preview': str(PREVIEW), 'ui_preview': str(UI_SCREEN)}, indent=2) + '\n', encoding='utf-8')
    log(f'[Layer analysis] PASS: {len(checks)} checks', 7, 7)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
