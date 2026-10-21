"""All open layers to one workbook, one sheet per layer. Workflow: README.md."""
from pathlib import Path
from app_identity import APP_NAME
import json
import re
import textwrap
import uuid

from attack_matrix import ALL_PLATFORMS
from matrix_exports import matrix_scene, cell_label, cell_note
from navigator_visuals import foreground_color
from matrix_display import marker_colors


def sheet_title(name, used):
    title=re.sub(r'[\x00-\x1f\\/*?:\[\]]','_',name).strip().strip("'") or 'Layer'
    def truncate(value,limit):
        return value.encode('utf-16-le')[:limit*2].decode('utf-16-le',errors='ignore')
    base=truncate(title,31).rstrip("'") or 'Layer'
    if base.casefold()=='history': base='History_'
    candidate=base
    suffix=2
    while candidate.casefold() in used:
        ending=f' ({suffix})'
        candidate=truncate(base,31-len(ending))+ending
        suffix+=1
    used.add(candidate.casefold())
    return candidate


def write_layer_sheet(sheet, scene, notify):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    def value(row,col,content):
        if isinstance(content,str) and (len(content)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',content)):
            raise ValueError('Excel text is too long or contains unsupported control characters; use JSON for complete data.')
        cell=sheet.cell(row,col,content)
        if isinstance(content,str): cell.data_type='s'
        cell.alignment=Alignment(wrap_text=True,vertical='top')
        return cell
    value(1,1,scene['name'])
    value(2,1,scene['scope']+' · '+scene['platform'])
    count=max(1,len(scene['columns']))
    for row in (1,2):
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=count)
        sheet.row_dimensions[row].height=30
    sheet.freeze_panes='A4'
    border=Border(**{side:Side(style='thin',color='41484D') for side in ('left','right','top','bottom')})
    bottom=3
    for number,column in enumerate(scene['columns'],1):
        header=value(3,number,column.get('label',column['name']))
        header_bg=scene.get('tactic_background','')
        header.fill=PatternFill('solid',fgColor=header_bg[1:].upper() if header_bg else '253A4E')
        header.font=Font(color=foreground_color(header_bg)[1:].upper() if header_bg else 'FFFFFF',bold=True)
        sheet.column_dimensions[get_column_letter(number)].width=36
        for row,data in enumerate(column['cells'],4):
            cell=value(row,number,cell_label(data))
            cell.fill=PatternFill('solid',fgColor=data['background'][1:].upper())
            cell.font=Font(color=data['foreground'][1:].upper(),size=11)
            cell.border=border
            if marker_colors(data['style']):
                from copy import copy
                indicators=copy(border)
                for side,color in zip(('bottom','right','top'),marker_colors(data['style'])):
                    setattr(indicators,side,Side(style='medium',color=color[1:].upper()))
                cell.border=indicators
            note=cell_note(data)
            # Excel comments can otherwise be silently truncated by consumers.
            if len(note)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',note):
                raise ValueError('Excel annotation is too long or contains unsupported control characters; use JSON.')
            cell.comment=Comment(note,APP_NAME)
            sheet.row_dimensions[row].height=max(sheet.row_dimensions[row].height or 0,
                20+15*sum(max(1,len(textwrap.wrap(line,30))) for line in cell_label(data).splitlines()))
            bottom=max(bottom,row)
        notify('[Workbook 2/4] Rendering tactic columns',number,len(scene['columns']))
    if not scene['columns']:
        value(3,1,'No techniques match this layer platform/visibility filters.')
        sheet.column_dimensions['A'].width=60
    details=[('Description',scene['description']),('Platform',scene['platform']),('Sub-techniques','All matching children'),
             ('Layout','Labeled vertical matrix'),('Metadata JSON',json.dumps(scene['metadata'],ensure_ascii=False)),
             ('Links JSON',json.dumps(scene.get('links',[]),ensure_ascii=False)),
             ('Formula JSON',json.dumps(scene.get('formula',{}),ensure_ascii=False))]
    details += [('Legend: '+entry['label'],entry['color']) for entry in scene['legend']]
    for index,(label,content) in enumerate(details,bottom+3):
        value(index,1,label)
        value(index,2,content)
    sheet.row_dimensions[3].height=36


def export_all_layers_excel(documents,target,notify):
    from openpyxl import Workbook
    target=Path(target).expanduser().resolve()
    if target.suffix.casefold()!='.xlsx': target=target.with_name(target.name+'.xlsx')
    temporary=target.with_name(target.name+'.'+uuid.uuid4().hex+'.tmp')
    notify('Input files: None (immutable snapshots of every open layer)',0,0)
    notify(f'Output file: {target}',0,0)
    notify(f'Temporary output file: {temporary}',0,0)
    notify('[Workbook 1/4] Validating all open layer snapshots',0,4)
    if not documents:
        raise ValueError('Open at least one layer before exporting.')
    if len(documents)>50:
        raise ValueError('Export at most 50 open layers per workbook.')
    book=Workbook()
    book.remove(book.active)
    used=set()
    try:
        for number,document in enumerate(documents,1):
            layer=document.navigator_layer
            platforms=layer.get('filters',{}).get('platforms',ALL_PLATFORMS)
            notify(f'[Workbook 2/4] Layer {number}/{len(documents)}: {document.name}',number-1,len(documents))
            try:
                scene=matrix_scene(document,platform=platforms,notify=notify,allow_empty=True)
            except ValueError as error:
                raise ValueError(f'Layer {number} ({document.name}): {error}') from error
            sheet=book.create_sheet(sheet_title(document.name,used))
            write_layer_sheet(sheet,scene,notify)
        notify('[Workbook 3/4] Writing the complete workbook',2,4)
        book.save(temporary)
        notify('[Workbook 4/4] Publishing output atomically',3,4)
        temporary.replace(target)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    notify(f'[Workbook 4/4] Complete: {len(documents)} layers; {target}',4,4)
    return target
