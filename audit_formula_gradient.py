"""NAV-020 official automatic gradient policies, native colors and exports. Flow: README.md."""
from copy import deepcopy
from pathlib import Path
from hashlib import sha256
import json
import math
from xml.etree import ElementTree as ET
from PySide6 import QtCore,QtGui,QtSvg,QtWidgets
from openpyxl import load_workbook
from attack_catalog import load_catalog, external_id, tactic_phases
from attack_matrix import AttackMatrix,MATRIX_SLUG,MATRIX_STYLE
from audit_formula_scope import fixture,documents
from audit_formula_features import Harness,wait
from custom_sources import load_custom
from layer_analysis import calculate_layer,automatic_formula_gradient,NAVIGATOR_FORMULA_GRADIENT
from layer_analysis_dialogs import FormulaDialog
from localization import language_manager
from main import configure_application
from matrix_exports import export_matrix,matrix_scene,SVG_NS
from navigator_integration import document_layer,imported_document
from navigator_visuals import LayerAppearance,DEFAULT_GRADIENT
from studio_files import save_studio,open_studio,session_payload,write_session,read_session
from ttp_selection import selected_cells

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'/'attack'
REF=ROOT/'scripts'/'formula-reference'
OUT=ROOT/'artifacts'/'formula-gradient'
CACHE=OUT/'cache'
CHECKS=[]


def log(message,current=0,total=0):
    print(message+(f' ({current}/{total})' if total else ''),flush=True)


def check(value,label):
    if not value: raise AssertionError(label)
    CHECKS.append(label)


def bad(task,label):
    try: task()
    except ValueError: check(True,label)
    else: raise AssertionError(label)


def make_documents(catalog,sources):
    docs=documents(catalog,sources)
    for doc,source in zip(docs,sources):
        if 'gradient' in source: doc.navigator_layer['gradient']=deepcopy(source['gradient'])
    return docs


def normalized(gradient):
    return {'colors':[color.lower()+('ff' if len(color)==7 else '') for color in gradient['colors']],
            'minValue':gradient['minValue'],'maxValue':gradient['maxValue']}


def main():
    for path in (REF/'gradient-cases.json',REF/'gradient-golden.json',DATA/'latest.json'):
        log('Input file: '+str(path.resolve()))
    log('Input directory: '+str(DATA.resolve()))
    log('Output cache directory: '+str(CACHE.resolve()))
    for name in ('audit.json','gradient-bundle.json','navigator.studio.json','session.json',
                 'gradient.png','gradient-zh.png','matrix.png','binary.svg','binary.xlsx','binary-preview.png'):
        log('Output file: '+str((OUT/name).resolve()))
    OUT.mkdir(parents=True,exist_ok=True)
    pointers={p:p.read_bytes() for p in (DATA/'latest.json',DATA/'active.json') if p.exists()}
    log('[Gradient Audit 1/7] Verify pinned reference and isolated dataset',0,7)
    cases=json.loads((REF/'gradient-cases.json').read_text())
    golden=json.loads((REF/'gradient-golden.json').read_text())
    check(golden['inputSHA256']==sha256((REF/'gradient-cases.json').read_bytes()).hexdigest(),'exact gradient fixture hash')
    check(golden['versions']=={'mathjs':'15.2.0','typescript':'5.8.3','tinycolor2':'1.6.0','tinygradient':'2.0.1'},'official reference dependencies pinned')
    check(golden['upstream'][1]['sha256']=='4e48d69a52ffaa221a8aaeedba7c1b3d31948ff14e4af5e0fd15e4e9b0650ec3','official real Gradient source pinned')
    path=OUT/'gradient-bundle.json'
    path.write_text(json.dumps(fixture(cases['matrix']),indent=2)+'\n',encoding='utf-8')
    catalog=load_custom(CACHE,str(path),'1.0',log)['catalog']
    ids={external_id(obj):key for key,obj in catalog.records.items()}
    results={}
    log('[Gradient Audit 2/7] Compare configurations, per-tactic scores and inherited states',1,7)
    for index,(case,expected) in enumerate(zip(cases['cases'],golden['results'],strict=True),1):
        docs=make_documents(catalog,case['sources'])
        before=[doc.fingerprint() for doc in docs]
        result=calculate_layer(docs,case['expression'],case['name'],inheritance=case['inheritance'],mode='navigator')
        check(normalized(result.navigator_layer['gradient'])==expected['gradient'],case['name']+': official serialized gradient')
        appearance=LayerAppearance(result.navigator_layer)
        selected=set()
        for cell in expected['cells']:
            actual=appearance.cell(cell['techniqueID'],cell['tactic'])
            score=cell['score']
            check(actual['score'] is None if score is None else actual['score'] is not None and
                  math.isclose(actual['score'],score,rel_tol=1e-13,abs_tol=1e-14),case['name']+': score '+cell['techniqueID']+'/'+cell['tactic'])
            check(actual['enabled']==cell['enabled'],case['name']+': enabled '+cell['techniqueID']+'/'+cell['tactic'])
            if score is not None: selected.add((ids[cell['techniqueID']],cell['tactic']))
        check(selected_cells(result)==selected,case['name']+': gradient policy does not alter inclusion')
        check([doc.fingerprint() for doc in docs]==before,case['name']+': inputs unchanged')
        check(result.navigator_layer['x_ttp_studio_formula']['gradient_mode']=='navigator',case['name']+': effective automatic policy recorded')
        results[case['name']]=result
        log('[Gradient Audit 2/7] Official case',index,len(cases['cases']))
    log('[Gradient Audit 3/7] Independent policy, explicit precedence and bounded failures',2,7)
    binary_case=next(c for c in cases['cases'] if c['name']=='binary')
    docs=make_documents(catalog,binary_case['sources'])
    nav=results['binary']
    studio=calculate_layer(docs,'a','Studio',scope='matrix')
    check(studio.navigator_layer['gradient']=={'colors':docs[0].navigator_layer['gradient']['colors'],'minValue':0,'maxValue':1},'default Studio palette/range unchanged')
    forced=calculate_layer(docs,'a','Forced Studio',mode='navigator',gradient_mode='studio')
    check(forced.navigator_layer['gradient']==studio.navigator_layer['gradient'],'Studio gradient override with Navigator calculation')
    check(forced.navigator_layer['techniques']==nav.navigator_layer['techniques'],'gradient override leaves source-score calculation unchanged')
    independent=calculate_layer(docs,'a','Studio scores official scale',scope='matrix',gradient_mode='navigator')
    check(independent.navigator_layer['gradient']==nav.navigator_layer['gradient'],'official gradient can be used independently with Studio scope')
    for mode in ('studio','navigator'):
        source=calculate_layer(docs,'a','Exact inherited',mode='navigator',gradient_mode=mode,inheritance={'gradient':'b'})
        check(source.navigator_layer['gradient']==docs[1].navigator_layer['gradient'],'explicit gradient source overrides '+mode)
    default=deepcopy(NAVIGATOR_FORMULA_GRADIENT)
    source=deepcopy(docs[0].navigator_layer['gradient'])
    mutable=automatic_formula_gradient([2,4],source,'navigator')
    mutable['colors'][0]='#00000000'
    check(NAVIGATOR_FORMULA_GRADIENT==default and docs[0].navigator_layer['gradient']==source,'automatic gradients share no mutable palette/source state')
    for invalid in ('invalid',None,False):
        bad(lambda:calculate_layer(docs,'a','Invalid',mode='navigator',gradient_mode=invalid),'reject invalid gradient policy: '+str(invalid))
    before=[doc.fingerprint() for doc in docs]
    bad(lambda:calculate_layer(docs,'a/0','Invalid',mode='navigator'),'numeric failure publishes no gradient result')
    def cancel(message,current=0,total=0):
        if '[Formula 4/5]' in message: raise InterruptedError('cancelled')
    try: calculate_layer(docs,'a','Cancelled',notify=cancel,mode='navigator')
    except InterruptedError: check(True,'cancellation before final gradient/result publication')
    else: raise AssertionError('missing cancellation')
    check([doc.fingerprint() for doc in docs]==before,'success/failure/cancel preserve source state')
    for domain in ('enterprise-attack','mobile-attack','ics-attack'):
        official=load_catalog(DATA,lambda *args:None).for_domain(domain)
        key,obj=next((k,o) for k,o in official.records.items() if len(tactic_phases(o,domain)&official.tactics.keys())>=2)
        slugs=[s for s in official.tactic_order if s in tactic_phases(obj,domain)]
        source_doc=make_documents(official,[{'rows':[{'techniqueID':external_id(obj),'tactic':slugs[0],'score':0},
            {'techniqueID':external_id(obj),'tactic':slugs[1],'score':1}]}])
        result=calculate_layer(source_doc,'a',domain,mode='navigator')
        check(result.navigator_layer['gradient']==nav.navigator_layer['gradient'],domain+': official binary gradient')
    log('[Gradient Audit 4/7] Studio/Navigator/recovery round trips and historical result preservation',3,7)
    save_studio(nav,OUT/'navigator.studio.json',log,save_as=True)
    reopened=open_studio(CACHE,OUT/'navigator.studio.json',log)
    check(reopened.navigator_layer==nav.navigator_layer,'Studio retains gradients and policy provenance')
    check(imported_document(document_layer(nav),catalog,'enabled').navigator_layer['gradient']==nav.navigator_layer['gradient'],'Navigator gradient JSON round trip')
    write_session(session_payload([nav,studio,results['constant inherited gradient']],0,{}),OUT/'session.json',log)
    restored=read_session(CACHE,OUT/'session.json',log)['documents']
    check([d.navigator_layer['gradient'] for d in restored]==[d.navigator_layer['gradient'] for d in [nav,studio,results['constant inherited gradient']]],'all gradient policies and inheritance survive recovery')
    old=deepcopy(document_layer(forced))
    old['x_ttp_studio_formula'].pop('gradient_mode')
    imported=imported_document(old,catalog,'enabled')
    check(imported.navigator_layer['gradient']==old['gradient'] and 'gradient_mode' not in imported.navigator_layer['x_ttp_studio_formula'],'pre-NAV-020 Navigator results preserve historical Studio gradient without recomputation')
    log('[Gradient Audit 5/7] Native controls, palette opacity and matrix rendering',4,7)
    app=QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    configure_application(app)
    parent=Harness()
    dialog=FormulaDialog(docs,docs[0],parent)
    dialog.show()
    app.processEvents()
    check(dialog.gradient_mode.currentData()=='auto','new UI follows calculation mode by default')
    dialog.calculation_mode.setCurrentIndex(dialog.calculation_mode.findData('navigator'))
    dialog.gradient_mode.setCurrentIndex(dialog.gradient_mode.findData('studio'))
    dialog.inheritance['gradient'].setCurrentIndex(dialog.inheritance['gradient'].findData('b'))
    check(not dialog.gradient_mode.isEnabled(),'explicit gradient inheritance disables automatic control')
    dialog.inheritance['gradient'].setCurrentIndex(0)
    check(dialog.gradient_mode.isEnabled() and dialog.gradient_mode.currentData()=='studio','clearing inheritance restores previous policy choice')
    dialog.gradient_mode.setCurrentIndex(0)
    manager=language_manager()
    manager.set_language('zh_TW');manager.refresh_all();app.processEvents()
    check(manager.translate(dialog.gradient_mode.currentText())!=dialog.gradient_mode.currentText(),'automatic gradient choice translated')
    check(dialog.grab().save(str(OUT/'gradient-zh.png')),'Chinese gradient preview')
    manager.set_language('en')
    check(dialog.grab().save(str(OUT/'gradient.png')),'English gradient preview')
    dialog.expression.setText('a')
    dialog.calculate()
    check(not dialog.options_tabs.isEnabled(),'gradient controls frozen during background work')
    wait(lambda:dialog.result_document is not None and not dialog.worker.isRunning())
    check(dialog.result_document.navigator_layer['gradient']==nav.navigator_layer['gradient'],'native default produces official binary gradient')
    check(len(parent.installed)==1 and parent.opened,'native worker publishes one independent result')
    check(any('[Formula 5/5]' in message[0] for message in parent.messages),'five-stage formula progress retained')
    matrix=AttackMatrix()
    matrix.resize(900,700)
    matrix.show()
    matrix.set_catalog(catalog,nav.selected)
    matrix.set_layer(document_layer(nav),catalog)
    wait(lambda:not matrix.building)
    for slug,opacity,background in [('execution',0,'#222526'),('persistence',1,'#66b1ff')]:
        item=next(item for item in matrix.items_by_id[ids['T0001']] if item.data(0,MATRIX_SLUG)==slug)
        style=item.data(0,MATRIX_STYLE)
        check(style['opacity']==opacity and style['background']==background,'native binary opacity/background: '+slug)
        check(item.background(0).color().name()==background,'native item uses projected gradient: '+slug)
    check(matrix.grab().save(str(OUT/'matrix.png')),'native binary matrix preview')
    manual=LayerAppearance(results['manual color inheritance does not change scale'].navigator_layer)
    check(manual.cell('T0001','execution')['background']=='#ff00ff','manual inherited color overrides transparent zero fill')
    disabled=LayerAppearance(results['disabled inherited result still counts'].navigator_layer)
    check(disabled.cell('T0001','execution')['background']=='#303436','disabled result retains its disabled appearance')
    matrix.close();parent.close();app.processEvents()
    log('[Gradient Audit 6/7] Actual SVG and colored Excel with alpha compositing',5,7)
    scene=matrix_scene(nav)
    check(scene['formula']['gradient_mode']=='navigator','export scene preserves automatic policy')
    export_matrix(nav,OUT/'binary.svg','svg',log)
    export_matrix(nav,OUT/'binary.xlsx','xlsx',log)
    for name in ('binary.svg','binary.xlsx'): log('Input file: '+str((OUT/name).resolve()))
    tree=ET.parse(OUT/'binary.svg')
    rects=[r for r in tree.findall('.//{'+SVG_NS+'}rect') if r.get('fill') in ('#222526','#66b1ff')]
    check(any(r.get('fill')=='#222526' for r in rects) and any(r.get('fill')=='#66b1ff' for r in rects),'SVG contains composited zero/one cell colors')
    renderer=QtSvg.QSvgRenderer(str(OUT/'binary.svg'))
    image=QtGui.QImage(renderer.defaultSize(),QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtCore.Qt.GlobalColor.transparent)
    painter=QtGui.QPainter(image);renderer.render(painter);painter.end()
    blue=next(r for r in rects if r.get('fill')=='#66b1ff')
    check(image.pixelColor(int(blue.get('x'))+int(blue.get('width'))-5,int(blue.get('y'))+5).name()=='#66b1ff','actual Qt SVG pixel matches blue endpoint')
    check(image.save(str(OUT/'binary-preview.png')),'exported SVG raster preview')
    book=load_workbook(OUT/'binary.xlsx')
    detail=list(book['Techniques'].values)[1:]
    check(all(row[8]==('#222526' if row[6]==0 else '#66b1ff') for row in detail),'Excel detail score/colors match native compositing')
    check(any(cell.fill.fgColor.rgb=='0066B1FF' for row in book['Matrix'] for cell in row),'Excel matrix blue fill')
    book.close()
    check(all(p.read_bytes()==data for p,data in pointers.items()),'official dataset pointers unchanged')
    log('[Gradient Audit 7/7] Publish report',6,7)
    (OUT/'audit.json').write_text(json.dumps({'passed':True,'count':len(CHECKS),'reference_gradients':len(golden['results']),
        'reference_cells':sum(len(r['cells']) for r in golden['results']),'checks':CHECKS},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    log(f'[Gradient Audit] PASS: {len(CHECKS)} checks',7,7)


if __name__=='__main__': main()
