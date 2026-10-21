"""Actual SVG rendering and all-layer Excel verification. Workflow: README.md."""
from copy import deepcopy
from pathlib import Path
import json
import time
from unittest.mock import patch
from xml.etree import ElementTree as ET

from PySide6 import QtCore, QtGui, QtWidgets, QtSvg
from openpyxl import load_workbook
from attack_catalog import load_catalog, external_id
from attack_matrix import MATRIX_ID, MATRIX_SLUG
from layer_analysis import snapshot_document
from main import MainWindow, configure_application
from matrix_exports import matrix_scene, svg_content, export_matrix, SVG_NS
from svg_export import validate_svg_options, DEFAULT_SVG, UNITS
from workbook_exports import export_all_layers_excel, sheet_title
from ttp_documents import TTPDocument
from navigator_integration import update_annotation

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'/'attack'
OUT=ROOT/'artifacts'/'export-options'
NS={'s':SVG_NS}
CHECKS=[]


def log(message,current=0,total=0):
    print(message+(f' ({current}/{total})' if total else ''),flush=True)


def check(value,label):
    if not value: raise AssertionError(label)
    CHECKS.append(label)


def bad(task,label):
    try: task()
    except (ValueError,PermissionError,InterruptedError): check(True,label)
    else: raise AssertionError(label)


def wait(condition):
    end=time.monotonic()+60
    while not condition():
        if time.monotonic()>end: raise TimeoutError('Native export did not finish')
        QtWidgets.QApplication.processEvents()
        time.sleep(.01)
    QtWidgets.QApplication.processEvents()


def main():
    log(f'Input file: {(DATA/"latest.json").resolve()}')
    for name in ('audit.json','advanced.svg','advanced.png','all-layers.xlsx','svg-settings.png','svg-preview.png','workbook-dialog.png'):
        log(f'Output file: {(OUT/name).resolve()}')
    OUT.mkdir(parents=True,exist_ok=True)
    pointers={p:p.read_bytes() for p in (DATA/'latest.json',DATA/'active.json') if p.exists()}
    app=QtWidgets.QApplication([])
    configure_application(app)
    catalog=load_catalog(DATA,log)
    ids={external_id(obj):key for key,obj in catalog.records.items()}
    parent,child=ids['T1053'],ids['T1053.005']
    a=TTPDocument('=Research A',catalog,selected={child})
    update_annotation(a,child,'execution','Literal <tag> & 中文 evidence',42.5,'#abcdef')
    a.navigator_layer.update(description='Worked example for publication.',legendItems=[{'label':'Assessment evidence','color':'#abcdef'}])
    scene=matrix_scene(a,'Windows','Scheduled Task/Job')
    check('Worked example for publication.' in svg_content(scene).decode('utf-8'),
          'default SVG includes the layer description with its about header')
    log('[Export Audit 1/7] SVG validation, units, fonts and actual vector rendering',0,7)
    options={**DEFAULT_SVG,'width':1600,'height':1000,'header_height':150,'font':'serif','cell_border':'#ff00aa',
             'legend_undocked':True,'legend_x':1200,'legend_y':180,'legend_width':350,'legend_height':120}
    export_matrix(a,OUT/'advanced.svg','svg',log,'Windows','Scheduled Task/Job',svg_options=options)
    root=ET.parse(OUT/'advanced.svg').getroot()
    check(root.get('width')=='1600px' and root.get('height')=='1000px','SVG explicit dimensions and units')
    body=root.find('s:g[@id="matrix"]',NS)
    check(body is not None and '150' in body.get('transform'),'explicit header height separates matrix body')
    check(all(node.get('font-family')=='Georgia, serif' for node in root.findall('.//s:text',NS)),'font applies to header, legend and cells')
    check(all(node.get('stroke')=='#ff00aa' for node in body.findall('.//s:rect',NS)),'custom cell border color')
    check(root.find('s:g[@id="legend"]',NS) is not None,'undocked legend is a separate foreground group')
    renderer=QtSvg.QSvgRenderer(str(OUT/'advanced.svg'))
    check(renderer.isValid() and renderer.defaultSize()==QtCore.QSize(1600,1000),'Qt renderer accepts SVG with exact canvas size')
    image=QtGui.QImage(renderer.defaultSize(),QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtCore.Qt.GlobalColor.transparent)
    painter=QtGui.QPainter(image)
    renderer.render(painter)
    painter.end()
    # Check a manual-color cell through the body's transform, including fit-to-canvas scaling.
    scale=float(body.get('transform').split('scale(')[1].split(')')[0])
    rect=next(node for node in body.findall('.//s:rect',NS) if node.get('fill')=='#abcdef')
    x=(float(rect.get('x'))+float(rect.get('width'))*.8)*scale
    y=150+(float(rect.get('y'))-100+float(rect.get('height'))*.8)*scale
    check(image.pixelColor(int(x),int(y)).name()=='#abcdef','actual SVG pixels preserve technique color through scaling')
    check(image.save(str(OUT/'advanced.png')),'SVG raster preview saved')
    for unit,width,height in (('cm',20,25),('in',8,10),('px',900,1200)):
        payload=svg_content(scene,{**DEFAULT_SVG,'unit':unit,'width':width,'height':height,'legend_width':300/UNITS[unit]})
        xml=ET.fromstring(payload)
        check(xml.get('width')==f'{width:g}'+unit and xml.get('height')==f'{height:g}'+unit,unit+': physical dimensions')
        rendered=QtSvg.QSvgRenderer(QtCore.QByteArray(payload))
        # Qt's default raster size uses a legacy 90 DPI for physical units;
        # SVG physical dimensions and the CSS 96-DPI viewBox remain explicit.
        check(rendered.isValid() and abs(rendered.viewBoxF().width()-width*UNITS[unit])<=.01,
              unit+': physical dimensions accepted with correct logical coordinate conversion')
    for mode in ('side','flat','mini'):
        current={**scene,'layout':{'layout':mode}}
        check(QtSvg.QSvgRenderer(QtCore.QByteArray(svg_content(current,options))).isValid(),mode+': advanced vector layout valid')
    hidden=ET.fromstring(svg_content(scene,{'show_header':False,'show_legend':False}))
    check(len(hidden.find('s:g[@id="header"]',NS))==0 and hidden.find('.//s:g[@id="legend"]',NS) is None,'header and legend can both be hidden')
    separate=ET.fromstring(svg_content(scene,{**options,'show_header':False}))
    check(separate.find('s:g[@id="legend"]',NS) is not None,'undocked legend survives hidden header')
    partial=ET.fromstring(svg_content(scene,{'show_about':False,'show_domain':False,'show_filters':False,'show_legend':False}))
    check(len(partial.find('s:g[@id="header"]',NS))==0,'header sub-sections toggle independently')
    docked=ET.fromstring(svg_content(scene,{'font':'monospace'}))
    check(docked.find('s:g[@id="header"]/s:g[@id="legend"]',NS) is not None,'docked legend follows header')
    invalid_xml=deepcopy(scene)
    invalid_xml['columns'][0]['cells'][0]['style']['comment']='invalid\x00XML'
    bad(lambda:svg_content(invalid_xml),'invalid XML in cell annotations fails visibly')
    bad(lambda:svg_content({**scene,'description':'invalid\x00header'},options),'invalid XML in advanced header fails visibly')
    for change in ({'width':-1},{'height':float('nan')},{'unit':'mm'},{'font':'evil'}, {'show_header':1},
                   {'cell_border':'red'},{'legend_width':10},{'subtechniques':'bad'},{'unknown':True}):
        bad(lambda change=change:validate_svg_options(change),'reject SVG option '+str(change))
    bad(lambda:svg_content(scene,{**options,'legend_x':1590}),'reject out-of-canvas legend')
    bad(lambda:svg_content(scene,{'height':50,'header_height':100}),'reject image with no matrix space')

    log('[Export Audit 2/7] Sub-technique modes and exact tactic expansion',1,7)
    flat=lambda projected:[cell for col in projected['columns'] for cell in col['cells']]
    none=matrix_scene(a,'Windows','Scheduled Task/Job',subtechniques='none')
    check(all(not cell['depth'] for cell in flat(none)),'show none removes all children')
    expanded=matrix_scene(a,'Windows','Scheduled Task/Job',subtechniques='expanded',expanded={(parent,'execution')})
    check(any(cell['depth'] for col in expanded['columns'] if col['slug']=='execution' for cell in col['cells']),'expanded execution includes children')
    check(not any(cell['depth'] for col in expanded['columns'] if col['slug']!='execution' for cell in col['cells']),'expanded mode does not leak across tactics')
    collapsed=matrix_scene(a,'Windows','Scheduled Task/Job',subtechniques='expanded',expanded=set())
    check(all(not cell['depth'] for cell in flat(collapsed)),'explicit empty expansion snapshot keeps all children collapsed')

    log('[Export Audit 3/7] One worksheet per layer, naming, independent domains and colors',2,7)
    b=snapshot_document(a)
    b.name='=research a'
    b.document_id='11111111-1111-4111-8111-111111111111'
    b.navigator_layer['name']=b.name
    b.navigator_layer['filters']={'platforms':[]}
    mobile=TTPDocument('Mobile: '+('Long name '*5),catalog.for_domain('mobile-attack'))
    mobile_key=next(iter(mobile.catalog.records))
    update_annotation(mobile,mobile_key,'','Mobile evidence',7,'#123456')
    ics=TTPDocument('ICS / matrix',catalog.for_domain('ics-attack'))
    docs=[a,b,mobile,ics]
    before=[doc.fingerprint() for doc in docs]
    export_all_layers_excel(docs,OUT/'all-layers.xlsx',log)
    book=load_workbook(OUT/'all-layers.xlsx')
    check(len(book.sheetnames)==4,'exactly one sheet per open layer with no extra sheets')
    check(book.sheetnames[:2]==['=Research A','=research a (2)'],'case-insensitive duplicate worksheet names deduplicated')
    check(all(len(name.encode('utf-16-le'))//2<=31 and not any(c in name for c in '[]:*?/\\') for name in book.sheetnames),'worksheet names satisfy Excel limits')
    check([sheet.cell(1,1).value for sheet in book]==[doc.name for doc in docs],'full names and tab order survive sanitization')
    check(all(cell.data_type!='f' for sheet in book for row in sheet for cell in row),'names and notes remain literal Excel text')
    check(all(sheet.freeze_panes=='A4' for sheet in book),'all sheets freeze matrix headers')
    found=[cell for row in book.worksheets[0] for cell in row if cell.comment and 'Literal <tag>' in cell.comment.text]
    check(found and found[0].fill.fgColor.rgb.endswith('ABCDEF'),'individual technique fills and literal comments retained')
    check('42.5' in found[0].value and 'Enabled: True' in found[0].comment.text,'scores and independent state included in cells/comments')
    check('No techniques' in book.worksheets[1].cell(3,1).value,'explicit empty platform filters retain a labeled empty sheet')
    check('Mobile' in book.worksheets[2].cell(2,1).value and 'ICS' in book.worksheets[3].cell(2,1).value,'mixed domains keep distinct scope headers')
    check(any(cell.fill.fgColor.rgb and str(cell.fill.fgColor.rgb).endswith('123456') for row in book.worksheets[2] for cell in row),'Mobile sheet keeps its own colors')
    book.close()
    used=set()
    for name in ('History',"'",'😀'*30,'A'*31,'a'*31,'a/b:c?d[e]','A'*30+"'"+'suffix'):
        title=sheet_title(name,used)
        check(bool(title) and len(title.encode('utf-16-le'))//2<=31 and title.casefold()!='history' and not title.endswith("'"),'safe sheet title: '+name[:10])
    check([doc.fingerprint() for doc in docs]==before,'multi-layer export preserves all source state')

    log('[Export Audit 4/7] Atomic failure, cancellation and complete workbook publication',3,7)
    protected=(OUT/'all-layers.xlsx').read_bytes()
    with patch.object(Path,'replace',side_effect=PermissionError('audit publication blocked')):
        bad(lambda:export_all_layers_excel(docs,OUT/'all-layers.xlsx',log),'workbook publication failure reported')
    def cancel(message,current=0,total=0):
        if '[Workbook 3/4]' in message: raise InterruptedError('audit cancelled')
    bad(lambda:export_all_layers_excel(docs,OUT/'all-layers.xlsx',cancel),'workbook cancellation stops before writing')
    oversized=snapshot_document(a)
    oversized.navigator_layer['description']='x'*32768
    bad(lambda:export_all_layers_excel([a,oversized],OUT/'all-layers.xlsx',log),'invalid later layer aborts entire workbook')
    bad(lambda:export_all_layers_excel([],OUT/'all-layers.xlsx',log),'empty layer list rejected')
    check((OUT/'all-layers.xlsx').read_bytes()==protected and not list(OUT.glob('*.tmp')),'failures/cancellation retain original workbook and remove temp files')
    svg_before=(OUT/'advanced.svg').read_bytes()
    with patch.object(Path,'replace',side_effect=PermissionError('audit publication blocked')):
        bad(lambda:export_matrix(a,OUT/'advanced.svg','svg',log,svg_options=options),'advanced SVG atomic publish failure reported')
    bad(lambda:export_matrix(a,OUT/'advanced.svg','svg',log,svg_options={'width':-1}),'invalid SVG rejected before publication')
    def cancel_svg(message,current=0,total=0):
        if '[Matrix Export 3/4]' in message: raise InterruptedError('audit SVG cancelled')
    bad(lambda:export_matrix(a,OUT/'advanced.svg','svg',cancel_svg,svg_options=options),'advanced SVG cancellation before publication')
    check((OUT/'advanced.svg').read_bytes()==svg_before and not list(OUT.glob('*.tmp')),'failed advanced SVG preserves old file')

    log('[Export Audit 5/7] Actual menu/dialog, preview, settings retry and workbook background tasks',4,7)
    window=MainWindow(DATA,DATA/'latest.json')
    window.show()
    wait(lambda:window.loaded_ok and not window.loader.isRunning())
    window.install_imported_document(snapshot_document(a))
    wait(lambda:window.loaded_ok)
    window.open_matrix()
    wait(lambda:not window.matrix_view.building)
    matrix=window.matrix_view
    matrix.set_expanded(False)
    check(not matrix.expansion_snapshot(),'native collapsed matrix captures empty expansion')
    matrix.set_expanded(True)
    check((parent,'execution') in matrix.expansion_snapshot(),'native expanded matrix captures technique/tactic pairs')
    for mode in ('side','mini','flat'):
        matrix.view_layout.setCurrentText(mode.title())
        wait(lambda:not matrix.building)
        matrix.set_expanded(True)
        check((parent,'execution') in matrix.expansion_snapshot(),mode+': native expansion capture')
    matrix.search.setText('Scheduled Task/Job')
    dialog=window.export_matrix_view('svg',OUT/'native.svg')
    settings=dialog.svg_settings
    check(settings is not None and settings.options() is None,'SVG settings preserve established defaults')
    settings.controls['width'].setValue(1600)
    settings.controls['height'].setValue(1000)
    settings.unit.setCurrentIndex(settings.unit.findData('cm'))
    check(abs(settings.controls['width'].value()*UNITS['cm']-1600)<.001,'changing units preserves physical size')
    settings.unit.setCurrentIndex(settings.unit.findData('px'))
    check(settings.controls['width'].value()==1600,'unit round-trip avoids near-integer display drift')
    settings.controls['font'].setCurrentIndex(settings.controls['font'].findData('serif'))
    settings.controls['subtechniques'].setCurrentIndex(settings.controls['subtechniques'].findData('expanded'))
    settings.controls['cell_border'].setText('invalid')
    dialog.preview_button.click()
    wait(lambda:not dialog.worker.isRunning() and dialog.preview_button.isEnabled())
    check('Failed:' in dialog.status.text() and dialog.output_path is None,'invalid native SVG settings remain editable without writing')
    settings.controls['cell_border'].setText('#ff00aa')
    dialog.report('Ready',1,1)
    check(dialog.grab().save(str(OUT/'svg-settings.png')),'advanced settings UI preview')
    preview_previous=(OUT/'native.svg').read_bytes() if (OUT/'native.svg').exists() else None
    dialog.preview_button.click()
    wait(lambda:not dialog.worker.isRunning() and dialog.preview_button.isEnabled())
    check(dialog.preview.renderer().isValid() and dialog.output_path is None,'background preview renders valid SVG without creating output')
    check(((OUT/'native.svg').read_bytes() if (OUT/'native.svg').exists() else None)==preview_previous,'preview never writes or overwrites the chosen destination')
    check(dialog.grab().save(str(OUT/'svg-preview.png')),'in-app SVG preview screenshot')
    dialog.export_button.click()
    wait(lambda:dialog.output_path is not None and not dialog.worker.isRunning())
    check(ET.parse(OUT/'native.svg').getroot().get('width').endswith('px'),'native SVG settings reach actual file export')
    dialog.reject()
    window.install_imported_document(snapshot_document(mobile))
    wait(lambda:window.loaded_ok)
    check(window.all_layers_excel_action.text()=='Export all open layers to Excel…','Navigator exposes all-layer Excel action')
    workbook=window.export_all_layers_workbook(OUT/'native.xlsx')
    check(workbook.layers_table.rowCount()==len(window.documents),'native workbook lists every open layer')
    check(workbook.grab().save(str(OUT/'workbook-dialog.png')),'workbook dialog screenshot')
    workbook.export_button.click()
    wait(lambda:workbook.output_path is not None and not workbook.worker.isRunning())
    native_book=load_workbook(OUT/'native.xlsx')
    check(len(native_book.sheetnames)==len(window.documents),'native workbook action exports every open layer once')
    native_book.close()
    workbook.reject()
    with patch.object(window,'confirm_studio_changes',return_value='discard'):
        window.close()
    log('[Export Audit 6/7] Dataset and source preservation',5,7)
    check(all(path.read_bytes()==data for path,data in pointers.items()),'official dataset pointers unchanged')
    check([doc.fingerprint() for doc in docs]==before,'preview/export/errors preserve source documents')
    log('[Export Audit 7/7] Publishing report',6,7)
    (OUT/'audit.json').write_text(json.dumps({'status':'PASS','checks':CHECKS},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    log(f'[Export Audit] PASS: {len(CHECKS)} checks',7,7)


if __name__=='__main__':
    main()
