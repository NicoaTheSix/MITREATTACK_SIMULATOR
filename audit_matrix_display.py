"""Verify labels, header colors, annotated expansion and indicators. Flow: README.md."""
from copy import deepcopy
from pathlib import Path
import json
import time
import uuid
from unittest.mock import patch
from xml.etree import ElementTree as ET

from PySide6 import QtCore, QtGui, QtWidgets, QtSvg
from openpyxl import load_workbook
from audit_advanced_navigator import fixture
from attack_catalog import external_id
from attack_matrix import MATRIX_ID, MATRIX_SLUG, MATRIX_STYLE
from custom_sources import load_custom
from layer_analysis import snapshot_document
from main import MainWindow, configure_application
from matrix_display import annotation_flags, display_label, marker_colors
from matrix_exports import matrix_scene, export_matrix, SVG_NS
from navigator_dialogs import LayerSettingsDialog
from navigator_integration import document_layer, update_annotation, validate_layer, imported_document, write_json_atomic
from navigator_visuals import LayerAppearance, OverlayAppearance
from studio_files import save_studio, open_studio, session_payload, write_session, read_session
from ttp_documents import TTPDocument, export_document_csv
from ttp_operations import batch_patch
from workbook_exports import export_all_layers_excel

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'/'attack'
OUT=ROOT/'artifacts'/'matrix-display'
CHECKS=[]
NS={'s':SVG_NS}


def log(message,current=0,total=0):
    print(message+(f' ({current}/{total})' if total else ''),flush=True)


def check(value,label):
    if not value: raise AssertionError(label)
    CHECKS.append(label)


def bad(task,label):
    try: task()
    except ValueError: check(True,label)
    else: raise AssertionError(label)


def wait(condition):
    end=time.monotonic()+60
    while not condition():
        if time.monotonic()>end: raise TimeoutError('Native matrix display task did not finish')
        QtWidgets.QApplication.processEvents()
        time.sleep(.01)
    QtWidgets.QApplication.processEvents()


def sample():
    payload=fixture()
    objects=payload['objects']
    ids={external_id(obj):obj for obj in objects if obj['type']=='attack-pattern'}
    root,child=ids['T1000'],ids['T1000.001']
    for obj in (root,child):
        obj['kill_chain_phases'].append({'kill_chain_name':'mitre-attack','phase_name':'initial-access'})
    for obj in objects:
        if obj['type']=='x-mitre-tactic':
            obj['external_references']=[{'source_name':'mitre-attack','external_id':'TA0002' if obj['x_mitre_shortname']=='execution' else 'TA0001'}]
    def add_child(identity,parent,name):
        obj=deepcopy(child)
        obj.update(id='attack-pattern--'+str(uuid.uuid5(uuid.NAMESPACE_DNS,identity)),name=name)
        obj['external_references'][0]['external_id']=identity
        objects.append(obj)
        objects.append({'type':'relationship','id':'relationship--'+str(uuid.uuid5(uuid.NAMESPACE_DNS,'relation '+identity)),
            'relationship_type':'subtechnique-of','source_ref':obj['id'],'target_ref':ids[parent]['id']})
        return obj
    sibling=add_child('T1000.002','T1000','Sibling')
    add_child('T1001.001','T1001','Zero score child')
    add_child('T1002.001','T1002','Unannotated child')
    for title,target,revoked in (('Evidence',child['id'],False),('Revoked note',sibling['id'],True)):
        objects.append({'type':'note','id':'note--'+str(uuid.uuid5(uuid.NAMESPACE_DNS,title)),
            'abstract':title,'content':'Read-only source note '+title,'object_refs':[target],
            'x_mitre_domains':['enterprise-attack'],'revoked':revoked})
    return payload


def item(matrix,key,slug='execution'):
    return next(value for value in matrix.items_by_id[key] if value.data(0,MATRIX_SLUG)==slug)


def main():
    log(f'Input directory: {DATA.resolve()}')
    log(f'Output directory: {OUT.resolve()}')
    for name in ('fixture.json','state.studio.json','state.navigator.json','state.csv','session.json','matrix.svg','matrix.xlsx','all-layers.xlsx','flat.png','side.png','mini.png','audit.json'):
        log(f'Output file: {(OUT/name).resolve()}')
    OUT.mkdir(parents=True,exist_ok=True)
    pointers={p:p.read_bytes() for p in (DATA/'latest.json',DATA/'active.json') if p.exists()}
    path=OUT/'fixture.json'
    path.write_text(json.dumps(sample(),ensure_ascii=False),encoding='utf-8')
    catalog=load_custom(DATA,str(path),'19.2',log)['catalog']
    ids={external_id(obj):key for key,obj in catalog.records.items()}
    root,child,sibling,beta,legacy=[ids[value] for value in ('T1000','T1000.001','T1000.002','T1001','T1002')]
    doc=TTPDocument('Display layer',catalog,selected={child})
    update_annotation(doc,child,'execution','Analyst note',7,'#113355',
        [{'label':'Evidence','url':'https://example.com/evidence'}],[{'name':'Analyst','value':'Team'}])
    update_annotation(doc,ids['T1001.001'],'execution','',0,'')
    update_annotation(doc,legacy,'','Parent annotation only',None,'')
    log('[Display Audit 1/7] Standard field validation and exact annotation scopes',0,7)
    for show_id in (False,True):
        for show_name in (False,True):
            layer={'layout':{'showID':show_id,'showName':show_name}}
            expected='  '.join(part for part,flag in (('T1000',show_id),('Alpha',show_name)) if flag)
            check(display_label('T1000','Alpha',layer)==expected,'independent label flags '+str((show_id,show_name)))
            check(display_label('T1000','Alpha',layer,True)=='','Mini suppresses label flags '+str((show_id,show_name)))
    check(not annotation_flags({'metadata':[{'name':'TTP Selection Studio selected','value':'true'}]})['annotated'],'selection markers are not annotations')
    check(not annotation_flags({'metadata':[{'divider':True}],'links':[{'divider':True}]})['annotated'],'dividers alone are not annotations')
    for row in ({'score':0},{'enabled':False},{'color':'#112233'}):
        state=annotation_flags(row)
        check(state['annotated'] and not marker_colors(state),'score/state/color count for expansion without note/link bars')
    style=LayerAppearance(document_layer(doc),catalog)
    cell=style.cell('T1000.001','execution')
    check(marker_colors(cell)==['#ffcc00','#4daafc'],'comments/metadata plus links create two distinct indicators')
    check('Evidence' in cell['sourceNotes'][0][0],'source STIX Notes indexed from exact catalog')
    check(style.cell('T1000.001','initial-access')['noteIndicator'] and not style.cell('T1000.001','initial-access')['linkIndicator'],'source notes span mapped tactics while scoped links do not')
    check(not style.cell('T1000.002','execution')['annotated'],'revoked source Note excluded')
    overlay=OverlayAppearance([('Evidence layer',document_layer(doc))],{},catalog)
    check(overlay.cell('T1000.001','execution')['linkIndicator'] and overlay.cell('T1000.001','execution')['annotationSources'],'overlay indicators expose source annotation context')
    for values in ({'layout':{'showID':'yes'}},{'layout':{'showName':1}}, {'showTacticRowBackground':1},{'tacticRowBackground':'red'}):
        bad(lambda values=values:validate_layer({'name':'Invalid','domain':catalog.domain,**values}),'invalid display fields rejected '+str(values))

    log('[Display Audit 2/7] Native label toggles, header color and annotated-only family expansion',1,7)
    app=QtWidgets.QApplication([])
    configure_application(app)
    window=MainWindow(DATA,DATA/'latest.json')
    window.resize(1500,950)
    window.show()
    wait(lambda:window.loaded_ok and not window.loader.isRunning())
    window.install_imported_document(doc)
    wait(lambda:window.loaded_ok)
    window.open_matrix()
    matrix=window.matrix_view
    wait(lambda:not matrix.building)
    check(item(matrix,child).text(0)=='Child','official default shows names without IDs')
    selected=set(doc.selected)
    matrix.show_ids.setChecked(True)
    matrix.show_names.setChecked(False)
    check(item(matrix,child).text(0)=='T1000.001','native ID-only label')
    matrix.show_ids.setChecked(False)
    check(item(matrix,child).text(0)=='' and 'T1000.001' in item(matrix,child).toolTip(0),'both labels off retains full tooltip identity')
    window.change_matrix_display({'showID':True,'showName':True,'showTacticRowBackground':True,'tacticRowBackground':'#f0c040'})
    check('T1000.001' in item(matrix,child).text(0) and 'Child' in item(matrix,child).text(0),'both labels on after atomic patch')
    heading=next(col['title'] for col in matrix.columns if col['slug']=='execution')
    check('TA0002' in heading.text() and '#f0c040' in heading.styleSheet() and '#000000' in heading.styleSheet(),'tactic header labels/background with readable foreground')
    token=doc.fingerprint()
    matrix.expand_annotated_button.click()
    expanded=matrix.expansion_snapshot()
    check((root,'execution') in expanded and (root,'initial-access') in expanded,'source Notes and annotations expand the correct family scopes')
    check((beta,'execution') in expanded,'zero-score child is annotated for expansion')
    check((legacy,'execution') not in expanded,'parent-only annotation does not expand unannotated children')
    check(doc.fingerprint()!=token and doc.selected==selected and doc.navigator_layer['layout']['expandedSubtechniques']=='annotated','annotated expansion persists policy and preserves selection')
    check('Read-only source note Evidence' in item(matrix,child).toolTip(0),'native source Notes appear as read-only tooltip content')
    app.processEvents()
    check(matrix.grab().save(str(OUT/'flat.png')),'Flat preview saved')

    log('[Display Audit 3/7] Painted indicators, clearing, Mini and Side behavior',2,7)
    update_annotation(doc,sibling,'execution','Temporary comment',None,'#113355',[{'label':'Case','url':'https://example.com/case'}])
    window.refresh_layer_appearance()
    target=item(matrix,sibling)
    tree=target.treeWidget()
    tree.scrollToItem(target)
    app.processEvents()
    rect=tree.visualItemRect(target)
    painted=tree.viewport().grab().toImage()
    check(painted.pixelColor(rect.left()+9,rect.bottom()-3).name()=='#ffcc00','Flat actually paints yellow note underline')
    check(painted.pixelColor(rect.left()+9,rect.bottom()-6).name()=='#4daafc','Flat actually paints blue link underline')
    batch_patch(doc,{sibling},{'comment':None,'links':None},'execution')
    window.refresh_layer_appearance()
    app.processEvents()
    target=item(matrix,sibling)
    rect=tree.visualItemRect(target)
    cleared=tree.viewport().grab().toImage()
    check(not marker_colors(target.data(0,MATRIX_STYLE)),'clear comment/links removes indicators')
    check(cleared.pixelColor(rect.left()+9,rect.bottom()-3).name() not in ('#ffcc00','#4daafc'),'clearing repaints without residual annotation line')
    for mode in ('Side','Mini'):
        matrix.view_layout.setCurrentText(mode)
        wait(lambda:not matrix.building)
        matrix.expand_annotated_button.click()
        check((root,'execution') in matrix.expansion_snapshot(),mode+': annotated expansion applies')
        column=next(c for c in matrix.columns if c['slug']=='execution')
        cell=next(c for f in column['view'].families for c in [f['root'],*f['children']] if c.key==child)
        check(marker_colors(cell.projection)==['#ffcc00','#4daafc'],mode+': note/link visual flags retained')
        image=cell.grab().toImage()
        check(image.pixelColor(8,cell.height()-5).name()=='#ffcc00',mode+': actual yellow indicator pixels')
        if mode=='Mini':
            check(not matrix.show_names.isEnabled() and not matrix.show_ids.isEnabled(),'Mini disables label controls without changing stored flags')
            check(column['title'].text()=='' and '#000000' in column['title'].styleSheet(),'Mini suppresses custom tactic background and heading text')
            check(doc.navigator_layer['layout']['showID'] and doc.navigator_layer['layout']['showName'],'Mini preserves label preferences')
        else:
            check('T1000.001' in cell.text() and 'Child' in cell.text(),'Side applies both label preferences')
        check(matrix.grab().save(str(OUT/(mode.lower()+'.png'))),mode+' screenshot')
    matrix.view_layout.setCurrentText('Flat')
    wait(lambda:not matrix.building)
    check(matrix.show_ids.isEnabled() and 'T1000.001' in item(matrix,child).text(0),'leaving Mini restores label preferences')
    settings=LayerSettingsDialog(doc,window)
    settings.show_names.setChecked(False)
    settings.show_ids.setChecked(True)
    settings.tactic_background.setText('invalid')
    settings.save()
    check(settings.result_layer is None and settings.error.text(),'invalid header color stays in reviewable settings draft')
    settings.tactic_background.setText('#225577')
    settings.save()
    check(settings.result()==QtWidgets.QDialog.DialogCode.Accepted and settings.result_layer['layout']['showID'],'Display tab saves a validated standard-field draft')

    log('[Display Audit 4/7] Per-layer switching, Studio/Navigator/session persistence',3,7)
    saved_fields=deepcopy(doc.navigator_layer)
    other=TTPDocument('Other labels',catalog)
    other.navigator_layer={'name':other.name,'domain':catalog.domain,'layout':{'showID':False,'showName':False}}
    window.install_imported_document(other)
    wait(lambda:window.loaded_ok and not matrix.building)
    check(not matrix.show_names.isChecked() and not matrix.show_ids.isChecked(),'switching layer restores its independent display flags')
    window.select_navigator_layer(doc.document_id)
    wait(lambda:window.loaded_ok and not matrix.building)
    check(matrix.show_names.isChecked() and matrix.show_ids.isChecked() and matrix.show_tactic_background.isChecked(),'returning layer restores all display controls')
    check(doc.navigator_layer==saved_fields,'layer switching does not rewrite display fields or annotations')
    save_studio(doc,OUT/'state.studio.json',log,save_as=True)
    reopened=open_studio(DATA,OUT/'state.studio.json',log)
    check(reopened.navigator_layer==doc.navigator_layer,'Studio reopen preserves display standard fields')
    check(LayerAppearance(document_layer(reopened),reopened.catalog).cell('T1000.001','execution')['sourceNotes'],'custom source Notes restored from verified cached Bundle')
    layer=document_layer(doc)
    write_json_atomic(layer,OUT/'state.navigator.json',log)
    imported=imported_document(layer,catalog)
    check(imported.navigator_layer['layout']==layer['layout'] and imported.navigator_layer['tacticRowBackground']=='#f0c040','Navigator exchange preserves standard display fields')
    write_session(session_payload([doc,other],0,{}),OUT/'session.json',log)
    recovered=read_session(DATA,OUT/'session.json',log)
    check(recovered['documents'][0].navigator_layer==doc.navigator_layer,'recovery preserves label and header settings')
    export_document_csv(doc,OUT/'state.csv',log)
    check('T1000.001' in (OUT/'state.csv').read_text(encoding='utf-8-sig'),'CSV keeps full TTP identity regardless of display choices')

    log('[Display Audit 5/7] SVG and colored Excel use the same labels/header/indicator projection',4,7)
    export_matrix(doc,OUT/'matrix.svg','svg',log)
    xml=ET.parse(OUT/'matrix.svg').getroot()
    texts=[node.text or '' for node in xml.findall('.//s:text',NS)]
    check(any('T1000.001' in text and 'Child' in text for text in texts),'SVG labels obey both flags')
    check(any(node.get('fill')=='#f0c040' for node in xml.findall('.//s:rect',NS)),'SVG tactic headers use saved background')
    colors={node.get('stroke') for node in xml.findall('.//s:line',NS)}
    check({'#ffcc00','#4daafc'}<=colors,'SVG includes both annotation indicator colors')
    renderer=QtSvg.QSvgRenderer(str(OUT/'matrix.svg'))
    check(renderer.isValid(),'actual Qt renderer accepts new SVG annotations')
    export_matrix(doc,OUT/'matrix.xlsx','xlsx',log)
    book=load_workbook(OUT/'matrix.xlsx')
    check(book['Matrix']['A3'].fill.fgColor.rgb.endswith('F0C040'),'single-layer Excel header color')
    cell=next(c for row in book['Matrix'] for c in row if c.comment and 'Analyst note' in c.comment.text)
    check(cell.border.bottom.color.rgb.endswith('FFCC00') and cell.border.right.color.rgb.endswith('4DAAFC'),'Excel note/link indicators use distinct borders')
    check('Read-only source note Evidence' in cell.comment.text,'Excel comment retains source STIX Note text')
    book.close()
    export_all_layers_excel([doc,other],OUT/'all-layers.xlsx',log)
    book=load_workbook(OUT/'all-layers.xlsx')
    check(book.worksheets[0]['A3'].fill.fgColor.rgb.endswith('F0C040'),'all-layer Excel applies each layer header color')
    check(book.worksheets[1]['A3'].value is None,'all-layer Excel honors labels disabled in another layer')
    book.close()
    for show_id,show_name in ((True,False),(False,True),(False,False)):
        probe=snapshot_document(doc)
        probe.navigator_layer['layout'].update(showID=show_id,showName=show_name)
        scene=matrix_scene(probe)
        displayed=next(c['display'] for col in scene['columns'] for c in col['cells'] if c['id']=='T1000.001')
        check(('T1000.001' in displayed)==show_id and ('Child' in displayed)==show_name,'export scene independent label choices '+str((show_id,show_name)))
    log('[Display Audit 6/7] Source, pointer and selection preservation',5,7)
    check(doc.selected==selected,'display/expansion never changes selected TTPs')
    check(all(p.read_bytes()==data for p,data in pointers.items()),'official dataset pointers unchanged')
    with patch.object(window,'confirm_studio_changes',return_value='discard'): window.close()
    log('[Display Audit 7/7] Publishing report',6,7)
    (OUT/'audit.json').write_text(json.dumps({'status':'PASS','checks':CHECKS},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    log(f'[Display Audit] PASS: {len(CHECKS)} checks',7,7)


if __name__=='__main__': main()
