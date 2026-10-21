"""Regression audit for native Navigator parity batch 1. Workflow: README.md."""
from copy import deepcopy
from pathlib import Path
import json
from itertools import product
from unittest.mock import patch
from PySide6 import QtCore, QtWidgets
from openpyxl import load_workbook
from audit_visibility_presets import fixture, wait
from attack_catalog import external_id
from attack_matrix import MATRIX_ID, SOURCE_SEARCH
from context_templates import CONTEXT_KEY, resolve_context, validate_context
from custom_sources import load_custom
from main import MainWindow, configure_application, ROLE_SEARCH
from matrix_display import marker_colors
from matrix_exports import matrix_scene, export_matrix
from navigator_dialogs import LayerSettingsDialog
from navigator_integration import document_layer, imported_document, validate_layer
from navigator_preferences import DISPLAY_KEY, SOURCE_KEY, DISPLAY_DEFAULTS
from navigator_visuals import LayerAppearance, OverlayAppearance, gradient_color, DEFAULT_GRADIENT
from studio_files import save_studio, open_studio, session_payload, write_session, read_session
from ttp_documents import TTPDocument
from workbook_exports import export_all_layers_excel

ROOT = Path(__file__).resolve().parent
OUT = ROOT/'artifacts'/'navigator-parity'
DATA = ROOT/'data'/'attack'
CHECKS = []


def log(text,current=0,total=0):
    print(text+(f' ({current}/{total})' if total else ''),flush=True)


def check(value,text):
    if not value: raise AssertionError(text)
    CHECKS.append(text)


def bad(callback,text):
    try: callback()
    except (ValueError,TypeError): check(True,text)
    else: raise AssertionError(text)


def main():
    log('Input directory: '+str(DATA.resolve()))
    log('Output directory: '+str(OUT.resolve()))
    for name in ('fixture.json','settings.ini','state.studio.json','session.json','matrix.svg','matrix.xlsx','all.xlsx','flat.png','side.png','mini.png','settings.png','audit.json'):
        log('Output file: '+str((OUT/name).resolve()))
    log('[Parity Audit 1/7] Preparing isolated source and validating extensions',0,7)
    OUT.mkdir(parents=True,exist_ok=True)
    bundle = fixture()
    for obj in bundle['objects']:
        if obj['type']=='x-mitre-tactic':
            identity='TA0002' if obj['x_mitre_shortname']=='execution' else 'TA0001'
            obj['external_references']=[{'source_name':'mitre-attack','external_id':identity,'url':'https://attack.mitre.org/tactics/'+identity+'/'}]
    for kind,identity in [('campaign','C1000'),('x-mitre-asset','A1000')]:
        bundle['objects'].append({'type':kind,'id':kind+'--a0000000-0000-4000-8000-000000000001',
            'name':'Research '+identity,'description':'Unique '+kind+' description with spaces.',
            'x_mitre_domains':['enterprise-attack'],'external_references':[{'source_name':'mitre-attack','external_id':identity}]})
    source = OUT/'fixture.json'
    source.write_text(json.dumps(bundle),encoding='utf-8')
    catalog = load_custom(DATA,str(source),'19.2',log)['catalog']
    ids = {external_id(obj):key for key,obj in catalog.records.items()}
    root,child = ids['T1000'],ids['T1000.001']
    layer = {'name':'Parity','domain':catalog.domain,'techniques':[
        {'techniqueID':'T1000','score':10},
        {'techniqueID':'T1000.001','score':40,'color':'#123456','comment':'Research note',
         'metadata':[{'name':'Owner','value':'Blue'}], 'links':[{'label':'Reference','url':'https://example.org/'}]}]}
    action = {'label':'Reference', 'url':'https://example.org/{{technique_attackID}}?t={{tactic_attackID}}',
              'subtechnique_url':'https://example.org/{{parent_technique_attackID}}/{{subtechnique_attackID_suffix}}?name={{subtechnique_name}}'}
    layer[CONTEXT_KEY] = [action]
    check(resolve_context(action,catalog,child,'execution').startswith('https://example.org/T1000/001?name='),'parent/sub-technique substitutions')
    check(resolve_context(action,catalog,root,'execution')=='https://example.org/T1000?t='+external_id(catalog.tactics['execution']),'technique/tactic substitutions')
    check(resolve_context({'label':'Child','url':'https://example.org/{{subtechnique_attackID}}'},catalog,root,'execution') is None,'missing sub-technique disables action')
    for template in ('javascript:alert(1)','file:///C:/test','https://user:secret@example.org/',
                     'https://{{technique_name}}.example.org/','https://example.org/{{unknown}}',
                     'https://example.org/{{technique_attackID}}}','https://example.org/\n'):
        bad(lambda template=template:validate_context({CONTEXT_KEY:[{'label':'Bad','url':template}]}),'reject unsafe or malformed template '+repr(template))
    for value in ([],['name','name'],['unknown'],[{}],False):
        bad(lambda value=value:validate_layer({**layer,SOURCE_KEY:value}),'reject invalid source fields '+repr(value))
    for value in ({'score_color':1},{'comment_color':'red'},{'unknown':True},[]):
        bad(lambda value=value:validate_layer({**layer,DISPLAY_KEY:value}),'reject invalid display setting '+repr(value))
    check(validate_layer(layer)[CONTEXT_KEY]==[action],'validated context templates retained losslessly')
    log('[Parity Audit 2/7] Independent color/indicator projection and overlays',1,7)
    original = deepcopy(layer)
    style = LayerAppearance(layer,catalog).cell('T1000.001','execution')
    check(style['background']=='#123456' and marker_colors(style)==['#ffcc00','#4daafc'],'legacy default colors/indicators preserved')
    layer[DISPLAY_KEY] = {'manual_color':False,'comment_color':'#112233','metadata_color':'#445566','link_color':'#778899'}
    style = LayerAppearance(layer,catalog).cell('T1000.001','execution')
    check(style['background']==gradient_color(40,DEFAULT_GRADIENT),'manual color hidden reveals original score')
    check(marker_colors(style)==['#112233','#445566','#778899'],'independent colored comment/metadata/link markers')
    layer[DISPLAY_KEY].update(score_color=False,aggregate_color=False,comment_indicator=False,metadata_indicator=False,link_indicator=False)
    style = LayerAppearance(layer,catalog).cell('T1000.001','execution')
    check(not style['background'] and not marker_colors(style) and style['annotated'],'hidden fills/indicators preserve classification')
    check(style['score']==40 and style['comment']=='Research note','hidden visuals preserve values')
    overlay = OverlayAppearance([('A',original)],layer,catalog).cell('T1000.001','execution')
    check(not marker_colors(overlay) and overlay['annotated'] and overlay['annotationSources'],'active overlay indicator prefs preserve source discovery')
    layer['layout'] = {'showAggregateScores':True,'aggregateFunction':'sum'}
    layer[DISPLAY_KEY] = {'manual_color':False,'score_color':False}
    style = LayerAppearance(layer,catalog).cell('T1000','execution')
    check(style['aggregateScore']==50 and style['background']==gradient_color(50,DEFAULT_GRADIENT),'aggregate color independent of original-score switch')
    layer[DISPLAY_KEY]['aggregate_color'] = False
    check(not LayerAppearance(layer,catalog).cell('T1000','execution')['background'],'all score fills disabled without deleting aggregate')
    check(original['techniques']==layer['techniques'],'visual projection never mutates annotations')
    for manual,score,aggregate in product((False,True),repeat=3):
        sample=deepcopy(layer)
        sample[DISPLAY_KEY]={'manual_color':manual,'score_color':score,'aggregate_color':aggregate}
        expected=gradient_color(50 if aggregate else 10,DEFAULT_GRADIENT) if aggregate or score else ''
        check(LayerAppearance(sample,catalog).cell('T1000','execution')['background']==expected,
              f'parent score switch combination {manual}/{score}/{aggregate}')
        expected='#123456' if manual else gradient_color(40,DEFAULT_GRADIENT) if score else ''
        check(LayerAppearance(sample,catalog).cell('T1000.001','execution')['background']==expected,
              f'child color switch combination {manual}/{score}/{aggregate}')
    for comment,metadata,link in product((False,True),repeat=3):
        sample=deepcopy(original)
        sample[DISPLAY_KEY]={'comment_indicator':comment,'metadata_indicator':metadata,'link_indicator':link,
                            'comment_color':'#112233','metadata_color':'#445566','link_color':'#778899'}
        style=LayerAppearance(sample,catalog).cell('T1000.001','execution')
        check(marker_colors(style)==[color for show,color in ((comment,'#112233'),(metadata,'#445566'),(link,'#778899')) if show],
              f'indicator switch combination {comment}/{metadata}/{link}')
    sample=deepcopy(original)
    sample['techniques'][1]['enabled']=False
    sample[DISPLAY_KEY]={'manual_color':False,'score_color':False}
    check(LayerAppearance(sample,catalog).cell('T1000.001','execution')['background']=='#303436','disabled state remains legible with fills off')
    previous_name=catalog.records[child]['name']
    catalog.records[child]['name']='Unicode 名称 / Path'
    check('unicode-%E5%90%8D%E7%A7%B0-%2F-path' in resolve_context(action,catalog,child,'execution'),'context names are slugged and URL encoded')
    catalog.records[child]['name']=previous_name
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    configure_application(app)
    w = MainWindow(DATA,settings_path=OUT/'settings.ini');w.resize(1800,1100);w.show()
    try:
        wait(lambda:w.loaded_ok and not w.loader.isRunning())
        document = TTPDocument('Parity',catalog,navigator_layer=deepcopy(original))
        w.install_imported_document(document);wait(lambda:w.loaded_ok)
        w.open_matrix();m=w.matrix_view;wait(lambda:not m.building)
        log('[Parity Audit 3/7] Search switches and document isolation',2,7)
        for category,needle in [('Campaigns','Unique campaign'),('Assets','Unique x-mitre-asset')]:
            w.switch_category(category)
            item=next(item for item in w.library_items[category] if '1000' in item.text())
            check(needle in item.data(ROLE_SEARCH),category+' library includes Description')
            w.search.setText(needle);w.apply_search()
            check(w.library_proxy.mapFromSource(item.index()).isValid(),category+' actual Description filter matches')
        m.source_kind.setCurrentText('Campaign')
        check(any('Unique campaign' in m.source_picker.itemData(i,SOURCE_SEARCH) for i in range(m.source_picker.count())), 'Campaign matrix completion includes Description')
        w.switch_category('Groups')
        check(w.source_search_fields_button.isVisible() and w.ttp_search_fields_button.isHidden(),'source and TTP controls separated')
        w.change_source_search_field('name',False);w.change_source_search_field('id',False)
        item = next(item for item in w.library_items['Groups'] if 'G1000' in item.text())
        check('description needle' in item.data(ROLE_SEARCH) and 'research g1000' not in item.data(ROLE_SEARCH).lower(),'sidebar index uses Description only')
        w.change_source_search_field('description',False)
        check(document.navigator_layer[SOURCE_KEY]==['description'],'last source field cannot be disabled')
        w.search.setText('description needle');w.apply_search()
        check(w.library_proxy.rowCount()>0,'source query with spaces remains usable')
        m.source_kind.setCurrentText('Group')
        idx=next(i for i in range(m.source_picker.count()) if 'G1000' in m.source_picker.itemText(i))
        check('description needle' in m.source_picker.itemData(idx,SOURCE_SEARCH),'matrix completer uses Description only')
        w.change_source_search_field('id',True)
        check(m.source_picker.currentText()=='' and not m.source_selection()[0],'changing fields clears stale source selection')
        w.change_source_search_field('description',False)
        check(item.data(ROLE_SEARCH)=='G1000','ID-only source search excludes Name and Description')
        w.change_source_search_field('name',True);w.change_source_search_field('id',False)
        check(item.data(ROLE_SEARCH)=='Research G1000','Name-only source search excludes ID projection and Description')
        w.change_source_search_field('id',True);w.change_source_search_field('description',True);w.change_source_search_field('name',False)
        w.switch_category('TTP')
        check(w.source_search_fields_button.isHidden() and w.ttp_search_fields_button.isVisible(),'TTP hides source controls')
        log('[Parity Audit 4/7] Saved drafts and annotations in all layouts',3,7)
        before = document.fingerprint()
        dialog = LayerSettingsDialog(document,w)
        dialog.effect_controls['comment_indicator'].setChecked(False)
        dialog.reject()
        check(document.fingerprint()==before,'cancel settings preserves complete document')
        dialog = LayerSettingsDialog(document,w)
        dialog.effect_controls['comment_color'].setText('bad')
        dialog.save()
        check(dialog.result_layer is None and dialog.error.text() and document.fingerprint()==before,'invalid settings do not commit')
        dialog.effect_controls['comment_color'].setText('#112233')
        dialog.effect_controls['metadata_color'].setText('#445566')
        dialog.effect_controls['link_color'].setText('#778899')
        dialog.save()
        check(dialog.result_layer is not None,'valid color/template draft commits')
        document.navigator_layer = dialog.result_layer
        document.navigator_layer[DISPLAY_KEY].update(comment_indicator=False,metadata_indicator=False,link_indicator=False)
        w.refresh_selection()
        for mode in ('Flat','Side','Mini'):
            m.view_layout.setCurrentText(mode);wait(lambda:not m.building)
            m.collapse_button.click()
            column=next(col for col in m.columns if col['slug']=='execution')
            if mode=='Flat':
                tree=column['tree'];parent=next(tree.topLevelItem(i) for i in range(tree.topLevelItemCount()) if tree.topLevelItem(i).data(0,MATRIX_ID)==root)
                check('[1 annotated]' in parent.text(0),'Flat collapsed annotated-child count')
                parent.setExpanded(True)
                check('[1 annotated]' not in parent.text(0),'Flat expanded label hides annotated count')
            else:
                family=next(f for f in column['view'].families if f['root'].key==root)
                check('1 annotated /' in family['root'].toolTip(),mode+' count remains discoverable with indicators off')
                if mode=='Side': check('[1 annotated]' in ' '.join(family['root'].text().split()),'Side collapsed count')
                family['toggle'].setChecked(True)
                if mode=='Side': check('[1 annotated]' not in ' '.join(family['root'].text().split()),'Side expanded label hides annotated count')
            check(m.grab().save(str(OUT/(mode.lower()+'.png'))),mode+' screenshot')
        settings=LayerSettingsDialog(document,w)
        settings.show()
        settings.context_editor.table.setCurrentCell(0,0)
        check(settings.grab().save(str(OUT/'settings.png')),'settings screenshot')
        settings.reject()
        document.navigator_layer['legendItems']=[{'label':'Research','color':'#123456'}]
        document.navigator_layer[DISPLAY_KEY]['legend_visible']=False
        w.refresh_selection()
        check(m.legend.isHidden(),'legend hides independently from technique colors')
        check(m.appearance.cell('T1000.001','execution')['background']=='#123456','hidden legend preserves technique fills')
        settings=LayerSettingsDialog(document,w)
        settings.clear_legend_button.click()
        check(settings.legend.rowCount()==0,'clear legend operates on whole draft')
        settings.reject()
        check(len(document.navigator_layer['legendItems'])==1,'cancel restores cleared legend draft')
        log('[Parity Audit 5/7] Actual context actions and safe URL dispatch',4,7)
        m.technique_menu(child,'execution',m,QtCore.QPoint(0,0))
        menu=m.cell_menu
        actions={a.text():a for a in menu.actions()}
        check('Select this technique' in actions and 'View technique details' in actions,'context commands target technique')
        seen=[]
        m.mouse_selection_requested.connect(lambda *args:seen.append(args))
        actions['Select only this technique'].trigger()
        check(seen[-1]==(child,'execution','replace'),'context command dispatches existing scoped preference path')
        custom=actions['Custom actions'].menu().actions()[0]
        with patch('PySide6.QtGui.QDesktopServices.openUrl',return_value=True) as opened:
            custom.trigger()
            check(opened.call_count==1 and '/T1000/001' in opened.call_args.args[0].toString(),'URL opened only through explicit action')
        menu.close()
        log('[Parity Audit 6/7] Studio/Navigator persistence and rendered exports',5,7)
        save_studio(document,OUT/'state.studio.json',log)
        reopened=open_studio(DATA,OUT/'state.studio.json',log)
        check(reopened.navigator_layer[DISPLAY_KEY]==document.navigator_layer[DISPLAY_KEY] and reopened.navigator_layer[SOURCE_KEY]==['id','description'],'Studio reopens native preferences')
        check(imported_document(document_layer(document),catalog).navigator_layer[CONTEXT_KEY]==[action],'Navigator round trip preserves extension templates')
        write_session(session_payload([document],0,{}),OUT/'session.json',log)
        recovered=read_session(DATA,OUT/'session.json',log)
        check(recovered['documents'][0].navigator_layer==document.navigator_layer,'session recovery preserves all native preferences and templates')
        document.navigator_layer[DISPLAY_KEY].update(comment_indicator=True,metadata_indicator=True,link_indicator=True)
        document.navigator_layer.setdefault('layout',{})['layout']='flat'
        scene=matrix_scene(document,notify=log)
        export_matrix(document,OUT/'matrix.svg','svg',log)
        export_matrix(document,OUT/'matrix.xlsx','xlsx',log)
        export_all_layers_excel([document],OUT/'all.xlsx',log)
        for path in (OUT/'matrix.xlsx',OUT/'all.xlsx'):
            workbook=load_workbook(path)
            colored=[cell for row in workbook.worksheets[0] for cell in row if cell.comment and 'Research note' in cell.comment.text]
            check(bool(colored),path.name+' includes annotated cell')
            cell=colored[0]
            check(cell.border.bottom.color.rgb=='00112233' and cell.border.right.color.rgb=='00445566' and cell.border.top.color.rgb=='00778899',path.name+' has three configured indicator colors')
            workbook.close()
        check('#112233' in (OUT/'matrix.svg').read_text(encoding='utf-8'),'SVG uses configured marker colors')
        log('[Parity Audit 7/7] Publishing verified report',6,7)
        (OUT/'audit.json').write_text(json.dumps({'passed':len(CHECKS),'checks':CHECKS},indent=2),encoding='utf-8')
        log(f'[Parity Audit] PASS: {len(CHECKS)} checks',7,7)
    finally:
        with patch.object(w,'confirm_studio_changes',return_value='discard'): w.close()
        QtWidgets.QApplication.processEvents()


if __name__=='__main__': main()
