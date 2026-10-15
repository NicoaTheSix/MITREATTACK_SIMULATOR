"""Full matrix SVG/XLSX exports with a shared color/score projection. Flow: README.md."""
from pathlib import Path
import json
import textwrap
import uuid
from xml.etree import ElementTree as ET

from attack_catalog import DOMAIN_LABELS, external_id, catalog_identity
from navigator_fields import sorted_families
from attack_matrix import ALL_PLATFORMS, matrix_columns
from navigator_integration import document_layer
from navigator_visuals import LayerAppearance, OverlayAppearance, display_color

SVG_NS = 'http://www.w3.org/2000/svg'


def matrix_scene(document, platform=ALL_PLATFORMS, query='', colors=True, sources=None, notify=lambda *args: None):
    catalog = document.catalog
    layer = document_layer(document)
    if sources is not None:
        compatible = [doc for doc in sources if catalog_identity(doc.catalog) == catalog_identity(catalog)]
        appearance = OverlayAppearance([(doc.name, document_layer(doc)) for doc in compatible], layer, catalog)
    else:
        appearance = LayerAppearance(layer, catalog)
    query = query.strip().casefold()
    columns = []
    matches = lambda key: query in (external_id(catalog.records[key]) + ' ' + catalog.records[key]['name']).casefold()
    projected = matrix_columns(catalog, platform)
    for number, (slug, families) in enumerate(projected, 1):
        families = sorted_families(catalog, slug, families, appearance, layer.get('sorting', 0))
        cells = []
        for family in families:
            key = family['key']
            root_matches = matches(key)
            children = [child for child in family['children'] if (root_matches or matches(child))
                        and not appearance.cell(external_id(catalog.records[child]), slug)['hidden']]
            root = appearance.cell(external_id(catalog.records[key]), slug)
            if not children and not (root_matches and not root['hidden']):
                continue
            for target, depth in [(key, 0), *[(child, 1) for child in children]]:
                record = catalog.records[target]
                style = appearance.cell(external_id(record), slug)
                parent = catalog.parents.get(target)
                cells.append({'key': target, 'id': external_id(record), 'name': record['name'], 'depth': depth,
                              'parent': external_id(catalog.records[parent]) if parent else '',
                              'selected': target in document.selected, 'style': style,
                              'background': style['background'] if colors and style['background'] else '#222526',
                              'foreground': style['foreground'] if colors and style['foreground'] else '#6caedf'})
        if cells:
            columns.append({'slug': slug, 'name': catalog.tactics[slug]['name'], 'cells': cells})
        notify('[Matrix Export 1/4] Projecting tactic columns', number, len(projected))
    if not columns:
        raise ValueError('No matrix cells match the current platform/search/visibility filters.')
    legend = list(appearance.layer.get('legendItems', [])) if colors else []
    if colors and sources is None and any('score' in row for row in appearance.rows.values()):
        gradient = appearance.gradient
        legend = [{'label': f"Gradient {gradient['minValue']:g}–{gradient['maxValue']:g}: {color}", 'color': color}
                  for color in gradient['colors']] + legend
    return {'name': document.name, 'scope': f'{DOMAIN_LABELS[catalog.domain]} v{catalog.version}',
            'platform': ', '.join(platform) if isinstance(platform, (list, tuple, set)) else platform,
            'query': query, 'overlay': sources is not None, 'colors': colors, 'columns': columns, 'legend': legend,
            'metadata': layer.get('metadata', []), 'description': layer.get('description', ''),
            'layout': layer.get('layout', {})}


def cell_label(cell):
    style = cell['style']
    label = ('> ' if cell['depth'] else '') + cell['id'] + '  ' + cell['name']
    if style['score'] is not None:
        label += f"\nScore: {style['score']:g}"
    if style.get('aggregateScore') is not None:
        label += f"\nAggregate: {style['aggregateScore']:g}"
    return label


def cell_note(cell):
    style = cell['style']
    parts = [cell_label(cell), 'Selected: ' + str(cell['selected']), 'Enabled: ' + str(style['enabled'])]
    if style['comment']:
        parts.append(style['comment'])
    for entry in style.get('metadata', []):
        if 'name' in entry and entry['name'] != 'TTP Selection Studio selected':
            parts.append(entry['name'] + ': ' + entry['value'])
    for entry in style.get('links', []):
        if 'label' in entry:
            parts.append(entry['label'] + ': ' + entry['url'])
    for source in style.get('sources', []):
        parts.append(f"{source['name']}: {source['color']} / score {source['score']} / aggregate {source.get('aggregateScore')}")
        if source['comment']:
            parts.append(source['comment'])
    return '\n'.join(parts)


def svg_content(scene):
    mode = scene['layout'].get('layout', 'flat')
    column_width = 464 if mode == 'side' else 232
    width = 36 + len(scene['columns']) * column_width
    legend_height = len(scene['legend']) * 22
    header = 100 + legend_height
    positions, tallest = [], header
    for column in scene['columns']:
        y, rows = header + 60, []
        for cell in column['cells']:
            lines = []
            for line in cell_label(cell).splitlines():
                lines.extend(textwrap.wrap(line, width=29 - cell['depth'] * 2, break_long_words=True) or [''])
            height = max(40, 14 + len(lines) * 15)
            rows.append((cell, y, height, lines, 0, 224))
            y += height
        if mode != 'flat':
            families = []
            for row in rows:
                if row[0]['depth'] == 0:
                    families.append([])
                families[-1].append(row)
            rows, y = [], header + 60
            for family in families:
                if mode == 'side':
                    root, children = family[0], family[1:]
                    family_height = max(root[2], sum(row[2] for row in children))
                    rows.append((root[0], y, family_height, root[3], 0, 222))
                    child_y = y
                    for child in children:
                        rows.append((child[0], child_y, child[2], child[3], 224, 232))
                        child_y += child[2]
                    y += family_height + 4
                else:
                    rows.append((family[0][0], y, 30, [], 0, 30))
                    for number, child in enumerate(family[1:]):
                        rows.append((child[0], y + (number // 5) * 32, 30, [], 38 + (number % 5) * 32, 30))
                    y += max(32, ((len(family) - 2) // 5 + 1) * 32) + 6
        positions.append((column, rows))
        tallest = max(tallest, y)
    ET.register_namespace('', SVG_NS)
    root = ET.Element(f'{{{SVG_NS}}}svg', width=str(width), height=str(tallest + 24), viewBox=f'0 0 {width} {tallest + 24}')
    def element(tag, **attrs):
        return ET.SubElement(root, f'{{{SVG_NS}}}{tag}', {key.replace('_', '-'): str(value) for key, value in attrs.items()})
    def text(x, y, value, color='#dddddd', size=12):
        item = element('text', x=x, y=y, fill=color, font_size=size, font_family='Arial, sans-serif')
        item.text = value
        return item
    element('rect', x=0, y=0, width=width, height=tallest + 24, fill='#1e1e1e')
    text(18, 28, scene['name'], size=20)
    text(18, 51, scene['scope'] + ' · ' + scene['platform'])
    text(18, 73, ('All compatible layers' if scene['overlay'] else 'Active layer') + ' · All matching sub-techniques expanded'
         + (' · Colors hidden' if not scene['colors'] else ''))
    for number, entry in enumerate(scene['legend']):
        x, y = 18, 88 + number * 22
        element('rect', x=x, y=y, width=12, height=12, fill=display_color(entry['color']))
        limit = max(16, (width - 75) // 6)
        label = entry['label'] if len(entry['label']) <= limit else entry['label'][:limit - 1] + '…'
        item = text(x + 18, y + 11, label, size=11)
        ET.SubElement(item, f'{{{SVG_NS}}}title').text = entry['label']
    for number, (column, rows) in enumerate(positions):
        x = 18 + number * column_width
        for line_no, line in enumerate(textwrap.wrap(column['name'], width=25)):
            text(x + 4, header + 18 + line_no * 16, line, '#6caedf', 13)
        for cell, y, height, lines, offset, cell_width in rows:
            rect = element('rect', x=x + offset, y=y, width=cell_width, height=height, fill=cell['background'],
                           stroke='#3794ff' if mode == 'mini' and cell['selected'] else '#41484d')
            title = ET.SubElement(rect, f'{{{SVG_NS}}}title')
            title.text = cell_note(cell)
            for line_no, line in enumerate(lines):
                text(x + offset + 7 + cell['depth'] * 8, y + 18 + line_no * 15, line, cell['foreground'])
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def save_excel(scene, path, notify):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    book = Workbook()
    matrix = book.active
    matrix.title = 'Matrix'
    details = book.create_sheet('Techniques')
    info = book.create_sheet('Layer Info')
    def value(sheet, row, col, content):
        cell = sheet.cell(row, col, content)
        if isinstance(content, str):
            if len(content) > 32767:
                raise ValueError('Excel cell text exceeds 32,767 characters; use JSON for full annotation text.')
            cell.data_type = 's'  # Names/comments are literal text, never Excel formulas.
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        return cell
    value(matrix, 1, 1, scene['name'])
    value(matrix, 2, 1, scene['scope'] + ' · ' + scene['platform'] + (' · All compatible layers' if scene['overlay'] else ''))
    matrix.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(scene['columns'])))
    matrix.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(scene['columns'])))
    matrix.row_dimensions[1].height = 26
    matrix.row_dimensions[2].height = 26
    matrix.freeze_panes = 'A4'
    border = Border(left=Side(style='thin', color='41484D'), right=Side(style='thin', color='41484D'),
                    top=Side(style='thin', color='41484D'), bottom=Side(style='thin', color='41484D'))
    headers = ['Tactic', 'Technique ID', 'Name', 'Parent ID', 'Selected', 'Enabled', 'Score', 'Aggregate Score', 'Color', 'Comment', 'Overlay Sources', 'Links JSON', 'Metadata JSON']
    for number, header in enumerate(headers, 1):
        value(details, 1, number, header)
    detail_row = 2
    try:
        for col_no, column in enumerate(scene['columns'], 1):
            cell = value(matrix, 3, col_no, column['name'])
            cell.fill = PatternFill('solid', fgColor='253A4E')
            cell.font = Font(color='FFFFFF', bold=True)
            matrix.column_dimensions[get_column_letter(col_no)].width = 36
            matrix.row_dimensions[3].height = 36
            for row_no, data in enumerate(column['cells'], 4):
                style = data['style']
                cell = value(matrix, row_no, col_no, cell_label(data))
                cell.fill = PatternFill('solid', fgColor=data['background'][1:].upper())
                cell.font = Font(color=data['foreground'][1:].upper(), size=11)
                cell.border = border
                cell.comment = Comment(cell_note(data), 'TTP Selection Studio')
                matrix.row_dimensions[row_no].height = max(matrix.row_dimensions[row_no].height or 0,
                    20 + 15 * sum(max(1, len(textwrap.wrap(line, 30))) for line in cell_label(data).splitlines()))
                parts = [column['name'], data['id'], data['name'], data['parent'], data['selected'], style['enabled'],
                         style['score'], style.get('aggregateScore'), data['background'], style['comment'],
                         '\n'.join(source['name'] for source in style.get('sources', [])),
                         json.dumps(style.get('links', []), ensure_ascii=False), json.dumps(style.get('metadata', []), ensure_ascii=False)]
                for index, part in enumerate(parts, 1):
                    value(details, detail_row, index, part)
                details.cell(detail_row, 2).fill = PatternFill('solid', fgColor=data['background'][1:].upper())
                details.cell(detail_row, 2).font = Font(color=data['foreground'][1:].upper())
                detail_row += 1
            notify('[Matrix Export 2/4] Rendering Excel columns', col_no, len(scene['columns']))
        details.freeze_panes = 'A2'
        details.auto_filter.ref = f'A1:M{detail_row - 1}'
        for index in range(1, 14):
            details.column_dimensions[get_column_letter(index)].width = 36 if index in (3, 10, 11) else 22
        metadata = [('Name', scene['name']), ('Domain/version', scene['scope']), ('Platform', scene['platform']),
                    ('Native / SVG layout', scene['layout'].get('layout', 'flat')),
                    ('Search', scene['query']), ('Colors shown', scene['colors']), ('Overlay', scene['overlay']),
                    ('Sub-techniques', 'All matching sub-techniques expanded'), ('Description', scene['description']),
                    ('Aggregate function', scene['layout'].get('aggregateFunction', 'average')),
                    ('Aggregate shown', scene['layout'].get('showAggregateScores', False)),
                    ('Count unscored as zero', scene['layout'].get('countUnscored', False))]
        metadata += [(entry.get('name', ''), entry.get('value', '')) for entry in scene['metadata']]
        metadata += [('Legend: ' + entry['label'], entry['color']) for entry in scene['legend']]
        for number, (key, content) in enumerate(metadata, 1):
            value(info, number, 1, key)
            value(info, number, 2, content)
        info.column_dimensions['A'].width = 30
        info.column_dimensions['B'].width = 80
        notify('[Matrix Export 3/4] Writing Excel workbook', 3, 4)
        book.save(path)
    finally:
        book.close()


def export_matrix(document, target, kind, notify, platform=ALL_PLATFORMS, query='', colors=True, sources=None):
    if kind not in ('svg', 'xlsx'):
        raise ValueError('Choose SVG or XLSX.')
    target = Path(target).expanduser().resolve()
    if target.suffix.lower() != '.' + kind:
        target = target.with_name(target.name + '.' + kind)
    temporary = target.with_name(target.name + '.' + uuid.uuid4().hex + '.tmp')
    notify('Input files: None (verified catalogs and layer snapshots in memory)', 0, 0)
    notify(f'Output file: {target}', 0, 0)
    notify(f'Temporary output file: {temporary}', 0, 0)
    notify('[Matrix Export 1/4] Validating layers and resolving filters', 0, 4)
    scene = matrix_scene(document, platform, query, colors, sources, notify)
    try:
        notify('[Matrix Export 2/4] Rendering complete matrix', 1, 4)
        if kind == 'svg':
            payload = svg_content(scene)
            notify('[Matrix Export 3/4] Writing SVG vector image', 2, 4)
            with temporary.open('xb') as stream:
                stream.write(payload)
        else:
            save_excel(scene, temporary, notify)
        notify('[Matrix Export 4/4] Publishing output atomically', 3, 4)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    notify(f'[Matrix Export 4/4] Complete: {target}', 4, 4)
    return target
