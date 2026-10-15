"""Native presentation, custom imports and annotation migration audit. Workflow: README.md."""
from copy import deepcopy
from dataclasses import replace
from hashlib import sha256
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import csv
import io
import json
import threading
import time
import uuid
from xml.etree import ElementTree as ET

from PySide6 import QtCore, QtWidgets, QtSvg, QtGui
from attack_catalog import external_id, load_catalog, local_path
from attack_matrix import AttackMatrix, matrix_columns, MATRIX_ID
from attack_versions import snapshot_pointer
from custom_sources import custom_catalog, load_custom, load_layer_url, read_source, collection_entries
from layer_analysis import calculate_layer
from layer_upgrade import upgrade_preview, migrate_layer
from main import MainWindow, configure_application
from matrix_exports import matrix_scene, export_matrix
from navigator_dialogs import AnnotationDialog, LayerSettingsDialog, ImportLayerDialog
from navigator_fields import set_presentation, sorted_families
from navigator_integration import document_layer, validate_layer, update_annotation, imported_document, compare_documents
from navigator_visuals import LayerAppearance
from source_dialogs import SourceDialog, UpgradeDialog
from ttp_documents import TTPDocument, document_payload, csv_content

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'data' / 'attack'
OUT = ROOT / 'artifacts' / 'advanced-navigator'
CHECKS = []


def log(message, current=0, total=0):
    print(message + (f' ({current}/{total})' if total else ''), flush=True)


def check(value, message):
    if not value:
        raise AssertionError(message)
    CHECKS.append(message)


def error(task, message):
    try:
        task()
    except (ValueError, TypeError, OSError):
        check(True, message)
    else:
        raise AssertionError(message)


def wait(condition, label):
    until = time.monotonic() + 90
    while not condition():
        if time.monotonic() > until:
            raise TimeoutError(label)
        QtWidgets.QApplication.processEvents()
        time.sleep(.01)
    QtWidgets.QApplication.processEvents()


def fixture(version='19.2', newer=False):
    domain = 'enterprise-attack'
    def obj(kind, suffix, **fields):
        return {'type': kind, 'id': kind + '--' + str(uuid.uuid5(uuid.NAMESPACE_DNS, kind + suffix)), 'name': suffix, 'x_mitre_domains': [domain], **fields}
    tactics = [obj('x-mitre-tactic', slug, x_mitre_shortname=slug, name=slug.title()) for slug in ['execution', 'initial-access']]
    matrix = obj('x-mitre-matrix', 'matrix', tactic_refs=[o['id'] for o in tactics])
    def technique(identity, name, platforms, scope='execution', **fields):
        return obj('attack-pattern', identity, name=name, x_mitre_platforms=platforms,
                   kill_chain_phases=[{'kill_chain_name': 'mitre-attack', 'phase_name': scope}],
                   external_references=[{'source_name': 'mitre-attack', 'external_id': identity}], **fields)
    records = [technique('T1000', 'Alpha changed' if newer else 'Alpha', ['Windows', 'Linux']),
               technique('T1000.001', 'Child', ['Windows'], 'initial-access' if newer else 'execution', x_mitre_is_subtechnique=True),
               technique('T1001', 'Beta', ['Linux']), technique('T1002', 'Legacy', ['Windows'], revoked=newer)]
    rel = {'type': 'relationship', 'id': 'relationship--' + str(uuid.uuid5(uuid.NAMESPACE_DNS, 'child')), 'relationship_type': 'subtechnique-of',
           'source_ref': records[1]['id'], 'target_ref': records[0]['id']}
    if newer:
        records.extend([technique('T1003', 'Replacement', ['Windows']), technique('T1004', 'Added', ['Linux'])])
        rels = [rel, {'type': 'relationship', 'id': 'relationship--' + str(uuid.uuid5(uuid.NAMESPACE_DNS, 'replacement')), 'relationship_type': 'revoked-by',
                     'source_ref': records[3]['id'], 'target_ref': records[4]['id']}]
    else:
        rels = [rel]
    collection = obj('x-mitre-collection', 'collection', x_mitre_version=version)
    return {'type': 'bundle', 'id': 'bundle--fixture', 'objects': [*tactics, matrix, collection, *records, *rels]}


def catalog(payload, source='https://example.org/custom.json'):
    data = json.dumps(payload).encode()
    return custom_catalog(payload, sha256(data).hexdigest(), source, '19.2', log)


def run():
    log(f'Input pointer: {(DATA / "latest.json").resolve()}')
    log(f'Output directory: {OUT.resolve()}')
    OUT.mkdir(parents=True, exist_ok=True)
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    configure_application(app)
    window = MainWindow(DATA, settings_path=OUT / 'settings.ini')
    window.resize(1450, 950)
    window.show()
    wait(lambda: window.loaded_ok, 'official startup')
    official = window.catalog
    pointers = {p: p.read_bytes() for p in [DATA / 'active.json', DATA / 'latest.json'] if p.exists()}
    original, target = catalog(fixture()), catalog(fixture('19.3', True))
    keys = {external_id(obj): key for key, obj in original.records.items()}
    document = TTPDocument('Advanced example', original, selected=set(keys.values()))
    update_annotation(document, keys['T1000'], '', 'Parent note', 8.5, '#f0c040',
                      [{'label': 'Case', 'url': 'https://example.org/case', 'extra': 'retained'}, {'divider': True}],
                      [{'name': 'Owner', 'value': 'Blue team', 'extra': 7}, {'divider': True}])
    update_annotation(document, keys['T1000.001'], 'execution', 'Scoped child note', 2, '#4455aa')
    update_annotation(document, keys['T1002'], '', 'Legacy observation', 4, '#aa4455')
    before = document.fingerprint()
    log('[Advanced Audit 1/6] Validating structured annotation editors')
    annotation = AnnotationDialog(document, keys['T1000'], window)
    check(annotation.links.table.rowCount() == 2 and annotation.metadata.table.rowCount() == 2, 'technique editor loads links, metadata and dividers')
    annotation.links.table.item(0, 0).setText('Updated case')
    annotation.metadata.table.item(0, 1).setText('Research team')
    annotation.save()
    check(annotation.result() == QtWidgets.QDialog.DialogCode.Accepted, 'technique structured edit accepted')
    check(document.fingerprint() == before, 'annotation draft leaves original unchanged')
    document.navigator_layer = annotation.result_layer
    row = next(row for row in document.navigator_layer['techniques'] if row['techniqueID'] == 'T1000')
    check(row['links'][0]['extra'] == 'retained' and row['metadata'][0]['extra'] == 7, 'entry extension fields preserved')
    check(row['links'][1] == {'divider': True}, 'divider preserved')
    invalid = AnnotationDialog(document, keys['T1000'], window)
    invalid.links.table.item(0, 1).setText('javascript:alert(1)')
    invalid.save()
    check(invalid.result() != QtWidgets.QDialog.DialogCode.Accepted and 'HTTP' in invalid.error.text(), 'unsafe edited link rejected atomically')
    settings = LayerSettingsDialog(document, window)
    settings.links.add({'label': 'Layer source', 'url': 'https://example.org/report'})
    settings.metadata.add({'name': 'Priority', 'value': 'High'})
    settings.save()
    check(settings.result() == QtWidgets.QDialog.DialogCode.Accepted, 'layer structured settings accepted')
    document.navigator_layer = settings.result_layer
    check(document_layer(document)['links'][0]['label'] == 'Layer source', 'layer links persist in Navigator JSON')
    check(document_payload(document)['attack']['dataset_sha256'] == original.dataset_id, 'Studio records custom dataset identity')
    check(next(csv.DictReader(io.StringIO(csv_content(document_payload(document)))))['attack_dataset_sha256'] == original.dataset_id, 'CSV records custom dataset identity')
    error(lambda: validate_layer({**document_layer(document), 'metadata': [{'name': 'bad', 'value': 42}]}), 'invalid metadata rejected')
    error(lambda: compare_documents(document, TTPDocument('official', official)), 'custom/official comparison blocked')
    other = TTPDocument('Other source', replace(original, dataset_id='another'), selected=set(keys.values()))
    error(lambda: calculate_layer([document, other], 'a+b', 'invalid'), 'cross-dataset arithmetic blocked')
    error(lambda: imported_document(document_layer(document), official), 'custom layer cannot import into official data')
    error(lambda: document.set_catalog(official), 'direct dataset switch blocked')

    log('[Advanced Audit 2/6] Exercising native layouts, sorting and platform unions')
    window.install_imported_document(document)
    wait(lambda: window.loaded_ok and window.catalog is original, 'custom document activation')
    window.open_matrix()
    matrix = window.matrix_view
    wait(lambda: not matrix.building, 'custom matrix')
    selected = set(document.selected)
    for mode in ('Side', 'Mini', 'Flat'):
        matrix.view_layout.setCurrentText(mode)
        wait(lambda: not matrix.building, mode + ' layout')
        check(document.navigator_layer['layout']['layout'] == mode.lower(), mode + ' saved in active layer')
        check(document.selected == selected, mode + ' preserves selection')
        if mode != 'Flat':
            check(all('view' in c for c in matrix.columns), mode + ' builds native family widgets')
            view = matrix.columns[0]['view']
            family = next(f for f in view.families if f['root'].key == keys['T1000'])
            check('Research team' in family['root'].toolTip(), mode + ' metadata visible in hover')
            matrix.set_expanded(True)
            check(family['toggle'].isChecked(), mode + ' expand all')
            matrix.set_expanded(False)
            check(not family['toggle'].isChecked(), mode + ' collapse all')
            family['root'].click()
            check(keys['T1000'] not in document.selected, mode + ' click updates actual document selection')
            family['root'].click()
        matrix.set_expanded(True)
        matrix.grab().save(str(OUT / (mode.lower() + '.png')))
    for sorting in range(4):
        matrix.sorting.setCurrentIndex(sorting)
        wait(lambda: not matrix.building, 'sorting')
        check(document.navigator_layer['sorting'] == sorting, 'sorting mode ' + str(sorting) + ' persists')
        slug, families = matrix.queue[0]
        expected = sorted_families(original, slug, dict(matrix_columns(original))[slug], matrix.appearance, sorting)
        check([r['key'] for r in families] == [r['key'] for r in expected], 'sorting mode ' + str(sorting) + ' projects correctly')
    order = [r['key'] for r in matrix.queue[0][1]]
    matrix.show_all_layers.setChecked(True)
    wait(lambda: not matrix.building, 'overlay score sort')
    check([r['key'] for r in matrix.queue[0][1]] == order, 'overlay color display retains active-layer score sorting')
    matrix.show_all_layers.setChecked(False)
    matrix.selected_platforms = ['Windows', 'Linux']
    matrix.presentation_changed()
    wait(lambda: not matrix.building, 'platform union')
    mapped = {key for _, rows in matrix_columns(original, ['Windows', 'Linux']) for r in rows for key in ([r['key']] if r['selectable'] else []) + r['children']}
    check(mapped == set(original.records), 'multi-platform OR union is complete')
    check(document.navigator_layer['filters']['platforms'] == ['Windows', 'Linux'], 'platform union persisted')
    matrix.selected_platforms = []
    matrix.presentation_changed()
    wait(lambda: not matrix.building, 'empty platform filter')
    check(not matrix.columns and matrix.empty.isVisible(), 'no selected platforms gives empty matrix')
    check(document.navigator_layer['filters']['platforms'] == [], 'empty platform selection persists in actual layer')
    matrix.platform.setCurrentText('Windows')
    wait(lambda: not matrix.building, 'single platform')
    check(document.navigator_layer['filters']['platforms'] == ['Windows'], 'single-platform shortcut replaces union')
    matrix.search.setText('T1000.001')
    check(any(c['frame'].isVisible() for c in matrix.columns), 'space-capable search works with platform filters')
    matrix.search.clear()
    for mode in ['Flat', 'Side', 'Mini']:
        matrix.view_layout.setCurrentText(mode)
        wait(lambda: not matrix.building, 'export layout')
        output = OUT / (mode.lower() + '.svg')
        export_matrix(document, output, 'svg', log, platform=['Windows', 'Linux'])
        renderer = QtSvg.QSvgRenderer(str(output))
        check(renderer.isValid(), mode + ' SVG is valid')
        xml = ET.parse(output)
        check('Legacy observation' in ''.join(xml.getroot().itertext()), mode + ' SVG retains technique annotation tooltip')
        image = QtGui.QImage(1100, 750, QtGui.QImage.Format.Format_ARGB32)
        image.fill(QtCore.Qt.GlobalColor.transparent)
        painter = QtGui.QPainter(image)
        renderer.render(painter)
        painter.end()
        image.save(str(OUT / (mode.lower() + '-svg.png')))
    export_matrix(document, OUT / 'advanced.xlsx', 'xlsx', log, platform=['Windows', 'Linux'])
    from openpyxl import load_workbook
    book = load_workbook(OUT / 'advanced.xlsx')
    check(book['Techniques'].cell(1, 12).value == 'Links JSON', 'Excel includes structured annotations')
    book.close()

    log('[Advanced Audit 3/6] URL imports, collections, bounds and custom validation')
    bundle = fixture()
    data = json.dumps(bundle).encode()
    bundle_path = OUT / 'custom-bundle.json'
    bundle_path.write_bytes(data)
    layer_data = json.dumps(document_layer(document)).encode()
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == '/bad':
                body = b'<html>not JSON</html>'
            elif self.path == '/bundle':
                body = data
            elif self.path == '/index':
                body = json.dumps({'collections': [{'name': 'Test collection', 'versions': [{'version': '19.2', 'url': '/bundle'}]}]}).encode()
            elif self.path == '/redirect':
                self.send_response(302)
                self.send_header('Location', 'file:///C:/Windows/win.ini')
                self.end_headers()
                return
            else:
                body = layer_data
            self.send_response(200)
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        def log_message(self, *args):
            pass
    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f'http://127.0.0.1:{server.server_port}'
    try:
        loaded = load_custom(OUT / 'cache', bundle_path, '19.2', log)
        check(loaded['catalog'].dataset_id == original.dataset_id and loaded['path'].read_bytes() == data, 'local custom bytes cache losslessly by content hash')
        remote = load_custom(OUT / 'cache', base + '/bundle', '19.2', log)
        check(remote['catalog'].dataset_id == original.dataset_id, 'HTTP custom import preserves dataset identity')
        index = load_custom(OUT / 'cache', base + '/index', '19.2', log)
        check(index['entries'][0]['url'] == base + '/bundle', 'collection relative URL resolves explicitly')
        url_layer = load_layer_url(OUT / 'cache', base + '/layer', log)
        check(url_layer['path'].read_bytes() == layer_data, 'URL layer JSON cached exactly')
        error(lambda: load_layer_url(OUT / 'cache', base + '/bad', log), 'HTML URL rejected')
        error(lambda: read_source(base + '/bundle', log, 32), 'URL size bound enforced')
        error(lambda: read_source(base + '/redirect', log), 'non-HTTP redirect blocked')
        error(lambda: load_layer_url(OUT / 'cache', 'javascript:alert(1)', log), 'unsupported URL scheme blocked')
        dialog = SourceDialog(OUT / 'cache', '19.2', window)
        dialog.kind.setCurrentIndex(1)
        dialog.location.setText(base + '/index')
        dialog.load()
        wait(lambda: dialog.worker and not dialog.worker.isRunning() and dialog.collection.count() == 1, 'collection UI preview')
        check(not dialog.commit_button.isEnabled() and dialog.result is None, 'collection index requires explicit Bundle choice')
        dialog.collection.setCurrentIndex(0)
        check(dialog.location.text() == base + '/bundle', 'collection UI fills selected source')
        dialog.load()
        wait(lambda: dialog.worker and not dialog.worker.isRunning() and dialog.result is not None, 'custom source UI')
        check(dialog.commit_button.isEnabled() and dialog.domain.currentData() == original.domain, 'custom preview enables exact domain import')
        received = []
        dialog.imported.connect(received.append)
        dialog.commit()
        check(len(received) == 1 and received[0].catalog.dataset_id == original.dataset_id, 'custom UI installs independent document')
        dialog.grab().save(str(OUT / 'custom-source.png'))
        layer_dialog = SourceDialog(OUT / 'cache', '19.2', window)
        layer_dialog.location.setText(base + '/layer')
        layer_dialog.load()
        wait(lambda: layer_dialog.worker and not layer_dialog.worker.isRunning() and layer_dialog.result, 'layer URL UI')
        paths = []
        layer_dialog.layer_ready.connect(paths.append)
        layer_dialog.commit()
        check(len(paths) == 1 and paths[0].is_absolute(), 'URL UI hands verified absolute cache path to existing importer')
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    broken = deepcopy(bundle)
    broken['objects'].append(deepcopy(broken['objects'][4]))
    broken['objects'][-1]['name'] = 'Conflict'
    error(lambda: catalog(broken), 'conflicting duplicate STIX version rejected')
    broken = deepcopy(bundle)
    broken['objects'][4]['kill_chain_phases'][0]['phase_name'] = 'missing'
    error(lambda: catalog(broken), 'unknown custom tactic rejected')
    broken = deepcopy(bundle)
    broken['objects'][-1]['target_ref'] = broken['objects'][-1]['source_ref']
    error(lambda: catalog(broken), 'custom parent cycle rejected')
    broken = deepcopy(bundle)
    broken['objects'].append({**broken['objects'][-1], 'id': 'relationship--' + str(uuid.uuid4()), 'target_ref': broken['objects'][6]['id']})
    error(lambda: catalog(broken), 'multiple active parents rejected instead of choosing a relationship silently')
    error(lambda: catalog({'type': 'bundle', 'objects': []}), 'empty custom Bundle rejected')
    latest = json.loads((DATA / 'latest.json').read_text(encoding='utf-8-sig'))
    real_payload, real_bytes, real_source = read_source(local_path(DATA, latest['enterprise_bundle']), log)
    real_custom = custom_catalog(real_payload, sha256(real_bytes).hexdigest(), real_source, latest['attack_version'], log)
    check(set(real_custom.records) == set(official.records), 'full official STIX Bundle loads as isolated custom catalog without losing techniques')
    del real_payload, real_bytes, real_custom
    for domain, pointer_key in [('mobile-attack', 'mobile_bundle'), ('ics-attack', 'ics_bundle')]:
        real_payload, real_bytes, real_source = read_source(local_path(DATA, latest[pointer_key]), log)
        custom = custom_catalog(real_payload, sha256(real_bytes).hexdigest(), real_source, latest['attack_version'], log)
        check(custom.domain == domain and set(custom.records) == set(official.for_domain(domain).records), domain + ' full custom Bundle preserves exact techniques')
        doc = TTPDocument('Custom ' + domain, custom, selected={next(iter(custom.records))})
        check(document_layer(doc)['domain'] == domain, domain + ' custom Navigator export supported')
    local_version = catalog(fixture('1.0'))
    doc = TTPDocument('Custom release 1.0', local_version, selected={next(iter(local_version.records))})
    check(document_layer(doc)['versions']['attack'] == '1.0', 'custom release 1.0 is not treated as historical official ATT&CK')
    check(calculate_layer([doc], 'a+1', 'Custom formula', missing='zero', notify=log).catalog is local_version, 'custom release 1.0 supports native formulas')
    imported = ImportLayerDialog(DATA, url_layer['path'], original, window)
    wait(lambda: imported.layers and imported.worker and not imported.worker.isRunning(), 'custom Navigator import resolution')
    check(imported.catalog.dataset_id == original.dataset_id and imported.import_button.isEnabled(), 'Navigator importer resolves already loaded custom digest')
    imported.reject()

    log('[Advanced Audit 4/6] Reviewed replacements, scope changes and collisions')
    before = document.fingerprint()
    preview = upgrade_preview(document, target, log)
    legacy = next(row for row in preview['rows'] if row['source_id'] == 'T1002')
    child = next(row for row in preview['rows'] if row['source_id'] == 'T1000.001' and row['source_scope'])
    check(legacy['target_id'] == 'T1003' and legacy['status'] == 'Replaced', 'explicit revoked-by suggests replacement')
    check(child['target_scope'] is None and not child['reviewed'], 'removed source tactic requires explicit remapping')
    check(any(change['field'] == 'kill_chain_phases' for change in child['changes']), 'upgrade shows exact before/after tactic mapping changes')
    check('T1004' in preview['added'], 'target additions reported separately')
    decisions = [{'action': row['action'], 'target_id': row['target_id'], 'target_scope': row['target_scope'],
                  'reviewed': row['reviewed']} for row in preview['rows']]
    error(lambda: migrate_layer(document, target, decisions, 'Migrated'), 'unreviewed migration blocked')
    for decision in decisions:
        decision['reviewed'] = True
    error(lambda: migrate_layer(document, target, decisions, 'Migrated'), 'invalid tactic cannot silently become global')
    for decision in decisions:
        if decision['target_scope'] is None:
            decision['target_scope'] = 'initial-access'
    migrated = migrate_layer(document, target, decisions, 'Migrated', log)
    check(document.fingerprint() == before, 'migration preserves source document exactly')
    migrated_row = next(row for row in migrated.navigator_layer['techniques'] if row['techniqueID'] == 'T1003')
    check(migrated_row['comment'] == 'Legacy observation' and migrated_row['score'] == 4 and migrated_row['color'] == '#aa4455', 'replacement preserves annotation fields')
    parent = next(row for row in migrated.navigator_layer['techniques'] if row['techniqueID'] == 'T1000')
    check(parent['links'][0]['extra'] == 'retained' and parent['metadata'][0]['extra'] == 7, 'upgrade preserves links and metadata extensions')
    target_ids = {external_id(obj): key for key, obj in target.records.items()}
    check(target_ids['T1003'] in migrated.selected and target_ids['T1004'] not in migrated.selected, 'selection migrates without auto-selecting target additions')
    check(migrated.catalog is target and migrated.navigator_layer['versions']['attack'] == '19.3', 'result uses exact target catalog and release')
    check(len(migrated.navigator_layer['x_ttp_studio_upgrade']['decisions']) == len(decisions), 'every migration decision retained in exported layer')
    check(migrated.migration_archives[0]['layer'] == document_layer(document), 'full source layer archived')
    check(not migrated.navigator_snapshots, 'custom source history cannot automatically restore into a different dataset')
    check(document_payload(migrated)['migration_archives'][0]['dataset_id'] == original.dataset_id, 'Studio JSON records dataset-scoped migration archive')
    conflict = deepcopy(decisions)
    for d in conflict:
        d.update(target_id='T1000', target_scope='')
    error(lambda: migrate_layer(document, target, conflict, 'Conflict'), 'conflicting destination scores/colors block migration')
    dropped = deepcopy(decisions)
    dropped[0]['action'] = 'Drop'
    result = migrate_layer(document, target, dropped, 'Dropped', log)
    check(result.navigator_layer['x_ttp_studio_upgrade']['decisions'][0]['action'] == 'Drop', 'explicit dropped row recorded')
    window.install_imported_document(TTPDocument('Target dataset', target))
    wait(lambda: window.loaded_ok and window.catalog is target, 'target custom document')
    upgrade = UpgradeDialog(DATA, document, window.documents, window)
    upgrade.target.setCurrentIndex(next(i for i in range(upgrade.target.count()) if upgrade.target.itemData(i)[0] == 'custom' and upgrade.target.itemData(i)[1] is target))
    upgrade.load_target()
    wait(lambda: upgrade.preview_data and upgrade.worker and not upgrade.worker.isRunning(), 'upgrade UI preview')
    check(upgrade.table.rowCount() == len(preview['rows']), 'review UI includes every source annotation')
    upgrade.table.setCurrentCell(next(i for i, row in enumerate(upgrade.preview_data['rows']) if row['source_id'] == 'T1000.001'), 0)
    check('kill_chain_phases' in upgrade.change_details.toPlainText() and 'Scoped child note' in upgrade.change_details.toPlainText(), 'native review exposes complete field differences and source annotation')
    for control in upgrade.controls:
        if control['scope'].currentData() is None:
            control['scope'].setCurrentIndex(control['scope'].findData('initial-access'))
    upgrade.mark_reviewed()
    upgrade.grab().save(str(OUT / 'upgrade-review.png'))
    results = []
    upgrade.imported.connect(results.append)
    upgrade.create()
    wait(lambda: results and upgrade.worker and not upgrade.worker.isRunning(), 'native migration creation')
    check(results[0].catalog is target and document.fingerprint() == before, 'native upgrade emits new document and preserves source')

    log('[Advanced Audit 5/6] Official historical upgrade and document isolation')
    historical = load_catalog(DATA, log, snapshot_pointer(DATA, '11.3'))
    key = next(key for key in historical.records if external_id(historical.records[key]) == 'T1053.005')
    old = TTPDocument('Historical annotations', historical, selected={key})
    update_annotation(old, key, '', 'Historical note', 6, '#556677', [{'label': 'Source', 'url': 'https://example.org/old'}], [{'name': 'Analyst', 'value': 'Research'}])
    preview = upgrade_preview(old, official, log)
    decision = [{'action': 'Copy', 'target_id': row['target_id'], 'target_scope': row['target_scope'], 'reviewed': True} for row in preview['rows']]
    migrated_official = migrate_layer(old, official, decision, 'Official upgraded', log)
    check(migrated_official.navigator_layer['techniques'][0]['comment'] == 'Historical note', 'real official 11.3 to latest annotation migration')
    check(not migrated_official.navigator_layer.get('customDataURL'), 'official result excludes custom source identity')
    mixed = migrate_layer(document, official, [{**d, 'action': 'Drop'} for d in decisions], 'Custom to official', log)
    check(not mixed.navigator_snapshots and mixed.migration_archives, 'custom-to-official migration archive excluded from automatic version restoration')
    window.new_document()
    wait(lambda: window.loaded_ok and not window.catalog.dataset_id, 'Ctrl+N equivalent from custom data')
    check(window.catalog.domain == 'enterprise-attack' and window.current_document.selected == set(), 'new file remains official Enterprise')
    check(all(path.read_bytes() == data for path, data in pointers.items()), 'official active/latest pointers unchanged by source imports and upgrades')
    for index in [window.document_index(document.document_id), len(window.documents) - 1]:
        window.activate_document(index)
        wait(lambda: window.loaded_ok, 'dataset tab switch')
    check(document.selected == selected, 'custom source selections survive tab switches')
    window.open_matrix()
    matrix = window.matrix_view
    wait(lambda: not matrix.building, 'full official matrix')
    projected = dict(matrix_columns(window.catalog))
    for mode in ['Side', 'Mini']:
        matrix.view_layout.setCurrentText(mode)
        wait(lambda: not matrix.building, 'full official ' + mode)
        check(len(matrix.columns) == len(projected), 'full official ' + mode + ' includes all tactic columns')
        for column in matrix.columns:
            check(sum(len(f['children']) for f in column['view'].families) == sum(len(f['children']) for f in projected[column['slug']]),
                  'full official ' + mode + '/' + column['slug'] + ' has all child cells')
        matrix.set_expanded(True)
        QtWidgets.QApplication.processEvents()
        if mode == 'Side':
            check(all(cell.height() >= cell.fontMetrics().height() + 10 for c in matrix.columns for f in c['view'].families for cell in f['children']),
                  'Side children retain readable heights in large families')
        matrix.grab().save(str(OUT / ('official-' + mode.lower() + '.png')))
    log('[Advanced Audit 6/6] Saving report')
    report = OUT / 'audit.json'
    log(f'Output file: {report.resolve()}')
    report.write_text(json.dumps({'passed': True, 'assertions': len(CHECKS), 'checks': CHECKS}, indent=2), encoding='utf-8')
    log(f'Advanced Navigator audit passed: {len(CHECKS)} assertions', 6, 6)
    window.closing = True
    window.close()


if __name__ == '__main__':
    run()
